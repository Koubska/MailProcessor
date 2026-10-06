"""Excel output: one workbook, appended to across runs.

Columns are found by their header name, never by position. Columns that are missing are added at the end;
nothing is moved or removed. So users can add their own columns and notes, and fields can be added,
removed or reordered without breaking an existing workbook. The workbook is in German, like its sheet names.
"""

from __future__ import annotations

import os
import uuid
from collections.abc import Callable, Mapping, Sequence
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from mailprocessor.columns import CONTENT_COLUMN, FIXED_DATA_COLUMNS, RECEIVED_COLUMN, TRANSFERRED_COLUMN
from mailprocessor.errors import SheetHeaderError, WorkbookLockedError, WorkbookUnreadableError

# The error sheet: one row per mail that currently fails. "Kennung" identifies the mail across runs; it is hidden.
KEY_COLUMN = "Kennung"
ERROR_COLUMNS = (
    "E-Mail",
    "Absender",
    "Betreff",
    "Eingegangen am",
    "Fehlende Felder",
    "Grund",
    "Geprüft am",
    KEY_COLUMN,
)
# English, technical layout of error sheets written by earlier versions; converted when the workbook is opened.
LEGACY_ERROR_COLUMNS = [
    "source_type",
    "source_location",
    "message_identity",
    "missing_columns",
    "error_reason",
    "processed_at",
]
# Texts of the "Grund" column; technical details follow in parentheses.
MISSING_FIELDS_REASON = "Pflichtfelder nicht gefunden"
UNREADABLE_REASON = "E-Mail konnte nicht gelesen werden"
INTERNAL_ERROR_REASON = "Interner Fehler beim Auslesen"
_LEGACY_MISSING_PREFIX = "Required fields missing:"

EXCEL_CELL_LIMIT = 32_767
_TRUNCATION_NOTE = "\n[… gekürzt: Excel erlaubt höchstens 32.767 Zeichen pro Zelle]"
DATE_FORMAT = "DD.MM.YYYY HH:MM"
# Column widths in characters. The mail text gets a fixed, wide column without wrapping,
# so each mail stays one row high; the full text shows when the cell is selected.
CONTENT_COLUMN_WIDTH = 80
DATE_COLUMN_WIDTH = 17
MIN_COLUMN_WIDTH = 10
MAX_COLUMN_WIDTH = 50
_WIDTH_SAMPLE_ROWS = 500
HEADER_FONT = Font(bold=True)


def error_key(source_type: str, source_location: str, message_identity: str) -> str:
    """Value of the hidden "Kennung" column: the same mail gets the same key in every run.

    Like the ledger, the key leaves out where the mail is (`source_location`): after moving the mail folder
    or renaming the mailbox, a mail that still fails keeps its one row instead of getting a second one.
    """
    return f"{source_type}|{message_identity}"


def _same_mail(stored: object, key: str) -> bool:
    """Whether a "Kennung" cell belongs to the mail with this key, also in the older "type|location|identity" form."""
    if not isinstance(stored, str):
        return False
    source_type, _, identity = key.partition("|")
    return stored == key or (stored.startswith(f"{source_type}|") and stored.endswith(f"|{identity}"))


@dataclass(frozen=True)
class ErrorEntry:
    """One row of the error sheet."""

    key: str  # see error_key
    name: str  # where users find the mail: file name or IMAP uid
    reason: str  # plain German explanation
    sender: str = ""
    subject: str = ""
    received: datetime | None = None
    missing: tuple[str, ...] = ()


def _cell_value(value: object) -> object:
    """Make untrusted text storable: drop control characters openpyxl rejects, respect Excel's cell limit."""
    if not isinstance(value, str):
        return value
    value = ILLEGAL_CHARACTERS_RE.sub("", value)
    if len(value) > EXCEL_CELL_LIMIT:
        value = value[: EXCEL_CELL_LIMIT - len(_TRUNCATION_NOTE)] + _TRUNCATION_NOTE
    return value


class _Sheet:
    """A worksheet whose columns are addressed by header name."""

    def __init__(self, sheet: Worksheet, wanted: Sequence[str]) -> None:
        self.sheet = sheet
        self.columns: dict[str, int] = {}
        last_header = 0
        for cell in sheet[1]:
            name = str(cell.value).strip() if cell.value is not None else ""
            if name:
                self.columns.setdefault(name, cell.column)
                last_header = cell.column
        last_row, last_column = _used_extent(sheet)
        if not self.columns and last_row > 1:
            raise SheetHeaderError(
                f"Sheet '{sheet.title}' has data but no header row; restore row 1 or choose a new output file."
            )
        # New rows go right after the last row with a value.
        self.next_row = last_row + 1
        # New columns go after everything that is used, also after notes in columns without a header.
        next_column = max(last_header, last_column) + 1
        for name in wanted:
            if name not in self.columns:
                sheet.cell(row=1, column=next_column, value=name)
                self.columns[name] = next_column
                next_column += 1

    def append(self, values: dict[str, object]) -> None:
        for name, value in values.items():
            cell = self.sheet.cell(row=self.next_row, column=self.columns[name], value=_cell_value(value))
            if isinstance(cell.value, str) and cell.value.startswith("="):
                # openpyxl turns any string starting with "=" into a live formula. Values come from
                # untrusted email content, so always store them as plain text (formula injection).
                cell.data_type = "s"
            elif isinstance(cell.value, datetime):
                cell.number_format = DATE_FORMAT
        self.next_row += 1

    def delete_rows_where(self, column: str, matches: Callable[[object], bool]) -> None:
        index = self.columns[column]
        for row_index in range(self.next_row - 1, 1, -1):
            if matches(self.sheet.cell(row=row_index, column=index).value):
                self.sheet.delete_rows(row_index)
                self.next_row -= 1

    def letter(self, column: str) -> str:
        return get_column_letter(self.columns[column])


def _used_extent(sheet: Worksheet) -> tuple[int, int]:
    """The last row and column below the header that hold a value (1 and 0 if there are none).

    Formatted but empty cells, e.g. borders or number formats set ahead in Excel, count for openpyxl's
    max_row/max_column; rows appended after them would end up far below the data.
    """
    last_row, last_column = 1, 0
    for row_index, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        used = [index for index, value in enumerate(row, start=1) if value is not None]
        if used:
            last_row, last_column = row_index, max(last_column, used[-1])
    return last_row, last_column


def _convert_legacy_error_sheet(sheet: Worksheet) -> None:
    """Rewrite an error sheet of an earlier version (English, technical) into the current German layout."""
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    sheet.delete_rows(1, sheet.max_row)
    for letter in list(sheet.column_dimensions):
        del sheet.column_dimensions[letter]
    converted = _Sheet(sheet, ERROR_COLUMNS)
    for source_type, location, identity, missing, reason, checked in rows:
        reason = str(reason or "")
        if reason.startswith(_LEGACY_MISSING_PREFIX):
            reason = MISSING_FIELDS_REASON
        try:
            checked_at = datetime.fromisoformat(str(checked)).astimezone().replace(tzinfo=None)
        except ValueError:
            checked_at = None
        converted.append(
            {
                "E-Mail": str(identity or "").removeprefix("file:"),
                "Fehlende Felder": missing or "",
                "Grund": reason,
                "Geprüft am": checked_at,
                KEY_COLUMN: error_key(str(source_type), str(location), str(identity)),
            }
        )


def _format_sheet(sheet: Worksheet, fixed_widths: dict[str, float]) -> None:
    """Bold header and readable column widths. Columns that already have a width
    (from an earlier run or set by the user in Excel) keep it."""
    for cell in sheet[1]:
        cell.font = HEADER_FONT
    # Excel stores one width for several columns as a range (e.g. B:D under "B"); a width added for a
    # column inside it would overlap, and Excel reports overlapping ranges as a damaged file.
    has_width: set[int] = set()
    for letter, dimension in sheet.column_dimensions.items():
        first = dimension.min or column_index_from_string(letter)
        has_width.update(range(first, max(dimension.max or first, first) + 1))
    sample_end = min(sheet.max_row, _WIDTH_SAMPLE_ROWS)
    for index, cells in enumerate(sheet.iter_cols(max_row=sample_end, values_only=True), start=1):
        letter = get_column_letter(index)
        if index in has_width:
            continue
        width = fixed_widths.get(letter)
        if width is None:
            lines = (line for value in cells if value is not None for line in str(value).splitlines())
            longest = max((len(line) for line in lines), default=0)
            width = min(max(longest + 2, MIN_COLUMN_WIDTH), MAX_COLUMN_WIDTH)
        sheet.column_dimensions[letter].width = width


def _open_workbook(path: Path) -> Workbook:
    try:
        return load_workbook(path)
    except OSError:
        raise  # e.g. no permission; handled like any other file error
    except Exception as exc:  # zipfile.BadZipFile, missing or broken parts inside the file, ...
        raise WorkbookUnreadableError(
            f"Cannot open {path} as an Excel workbook ({type(exc).__name__}). It may be damaged, not an .xlsx file, "
            "or protected with a password. Remove the password, or choose another output file."
        ) from None


def _locked_message(path: Path) -> str:
    return f"Cannot write {path}. Is it open in Excel? Close it and run again."


class ExcelOutput:
    """In-memory workbook that is written to disk once, atomically, via `save()`.

    `data_sheets` maps each data sheet to its field columns (one sheet, or one per profile).
    `now` is the time of the run: "Übertragen am" in the data sheets and "Geprüft am" in the error sheet.
    """

    def __init__(
        self,
        path: Path,
        data_sheets: Mapping[str, Sequence[str]],
        error_sheet: str,
        now: datetime | None = None,
    ) -> None:
        self.path = path
        self.now = now or datetime.now().replace(microsecond=0)
        self.is_new = not path.exists()

        if self.is_new:
            self.workbook = Workbook()
            self.workbook.active.title = next(iter(data_sheets))
        else:
            self.workbook = _open_workbook(path)
        for name in (*data_sheets, error_sheet):
            # Excel compares sheet names ignoring case, so "Daten" is the existing sheet "daten" (renamed),
            # not a new one; openpyxl would otherwise create "Daten1" and the lookup below would fail.
            existing = next((title for title in self.workbook.sheetnames if title.casefold() == name.casefold()), None)
            if existing is None:
                self.workbook.create_sheet(name)
            elif existing != name:
                # Via a temporary name: openpyxl counts the sheet's own old name as a duplicate ("Daten1").
                sheet = self.workbook[existing]
                sheet.title = f"~{uuid.uuid4().hex[:20]}"
                sheet.title = name
        errors = self.workbook[error_sheet]
        if [cell.value for cell in errors[1]] == LEGACY_ERROR_COLUMNS:
            _convert_legacy_error_sheet(errors)
        self.data = {
            name: _Sheet(self.workbook[name], [*columns, *FIXED_DATA_COLUMNS]) for name, columns in data_sheets.items()
        }
        self.errors = _Sheet(errors, ERROR_COLUMNS)

    def check_writable(self) -> None:
        """Fail early (before any processing) if the workbook is locked, e.g. open in Excel on Windows."""
        if self.is_new:
            return
        try:
            with self.path.open("r+b"):
                pass
        except PermissionError:
            raise WorkbookLockedError(_locked_message(self.path)) from None

    def append_data(
        self, values: dict[str, str], content: str, received: datetime | None = None, sheet: str | None = None
    ) -> None:
        """Add one row; `sheet` defaults to the first data sheet."""
        target = self.data[sheet] if sheet is not None else next(iter(self.data.values()))
        target.append({**values, RECEIVED_COLUMN: received, TRANSFERRED_COLUMN: self.now, CONTENT_COLUMN: content})

    def remove_errors_for(self, key: str) -> None:
        self.errors.delete_rows_where(KEY_COLUMN, lambda stored: _same_mail(stored, key))

    def upsert_error(self, entry: ErrorEntry) -> None:
        """Keep exactly one error row per mail, so retried failures don't pile up."""
        self.remove_errors_for(entry.key)
        self.errors.append(
            {
                "E-Mail": entry.name,
                "Absender": entry.sender,
                "Betreff": entry.subject,
                "Eingegangen am": entry.received,
                "Fehlende Felder": ", ".join(entry.missing),
                "Grund": entry.reason,
                "Geprüft am": self.now,
                KEY_COLUMN: entry.key,
            }
        )

    def save(self) -> None:
        formats = []
        for data in self.data.values():
            widths = {data.letter(name): DATE_COLUMN_WIDTH for name in (RECEIVED_COLUMN, TRANSFERRED_COLUMN)}
            widths[data.letter(CONTENT_COLUMN)] = CONTENT_COLUMN_WIDTH
            formats.append((data, widths))
        error_widths = {self.errors.letter(name): DATE_COLUMN_WIDTH for name in ("Eingegangen am", "Geprüft am")}
        for sheet, widths in (*formats, (self.errors, error_widths)):
            sheet.sheet.freeze_panes = "A2"
            sheet.sheet.auto_filter.ref = sheet.sheet.dimensions
            _format_sheet(sheet.sheet, widths)
        self.errors.sheet.column_dimensions[self.errors.letter(KEY_COLUMN)].hidden = True
        self.path.parent.mkdir(parents=True, exist_ok=True)
        temp_path = self.path.with_name(f".{self.path.name}.tmp")
        try:
            self.workbook.save(temp_path)
            os.replace(temp_path, self.path)
        except PermissionError:
            raise WorkbookLockedError(_locked_message(self.path)) from None
        finally:
            temp_path.unlink(missing_ok=True)
