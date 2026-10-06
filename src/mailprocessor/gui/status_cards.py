"""The status cards on the Start tab: is everything ready for a run?"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from mailprocessor.config import FieldRule
from mailprocessor.gui.config_files import setting_path
from mailprocessor.gui.form import FormValues, form_text
from mailprocessor.gui.preview import RulePreview, sample_files
from mailprocessor.gui.rule_inputs import split_labels
from mailprocessor.i18n import quote, t
from mailprocessor.parser import ParseResult


@dataclass(frozen=True)
class CardStatus:
    """One status card on the start page. ok: True = ready, False = needs attention, None = neutral."""

    ok: bool | None
    text: str


def mail_count_in_folder(values: FormValues, config_dir: Path) -> int | None:
    """Number of mail files in the configured folder, or None if the folder does not exist."""
    folder = setting_path(form_text(values, "eml_folder") or ".", config_dir)
    if not folder.is_dir():
        return None
    return len(sample_files(folder, form_text(values, "eml_glob") or "*.eml"))


def mails_status(values: FormValues, config_dir: Path, password_given: bool, lang: str) -> CardStatus:
    status = _source_status(values, config_dir, password_given, lang)
    if not status.ok:
        return status
    # Ready: say which mails count, so nobody wonders why some are not in the workbook.
    or_word = f" {t('card.mails.filter_or', lang)} "
    notes = [
        t(key, lang).format(entries=or_word.join(quote(entry, lang) for entry in entries))
        for key, entries in (
            ("card.mails.filter_subject", split_labels(form_text(values, "filter_subject"))),
            ("card.mails.filter_sender", split_labels(form_text(values, "filter_sender"))),
        )
        if entries
    ]
    return CardStatus(status.ok, " · ".join([status.text, *notes]))


def _source_status(values: FormValues, config_dir: Path, password_given: bool, lang: str) -> CardStatus:
    if form_text(values, "source_type") == "imap":
        host, user = form_text(values, "imap_host"), form_text(values, "imap_username")
        if not host or not user:
            return CardStatus(False, t("card.mails.imap_incomplete", lang))
        if not password_given:
            return CardStatus(False, t("card.mails.imap_password", lang).format(user=user))
        return CardStatus(True, t("card.mails.imap_ready", lang).format(user=user, host=host))
    folder_name = quote(Path(form_text(values, "eml_folder") or ".").name or form_text(values, "eml_folder"), lang)
    count = mail_count_in_folder(values, config_dir)
    if count is None:
        return CardStatus(False, t("card.mails.folder_missing", lang).format(folder=folder_name))
    if count == 0:
        return CardStatus(False, t("card.mails.folder_empty", lang).format(folder=folder_name))
    return CardStatus(True, t("card.mails.folder_ready", lang).format(folder=folder_name, count=count))


def fields_status(rules: list[FieldRule], previews: list[RulePreview] | None, lang: str) -> CardStatus:
    if not rules:
        return CardStatus(False, t("card.fields.none", lang))
    count = t("card.fields.count", lang).format(count=len(rules))
    if previews is None:
        return CardStatus(None, count)
    missing = [
        rule.column for rule, preview in zip(rules, previews, strict=True) if rule.required and not preview.found
    ]
    if missing:
        columns = ", ".join(quote(column, lang) for column in missing)
        return CardStatus(False, t("card.fields.missing", lang).format(count=count, columns=columns))
    return CardStatus(True, t("card.fields.ok", lang).format(count=count))


def profiles_status(profile_count: int, field_count: int, best: ParseResult | None, lang: str) -> CardStatus:
    """The fields card when there are several profiles: which profile the sample mail fits."""
    count = t("card.profiles.count", lang).format(profiles=profile_count, fields=field_count)
    if best is None:
        return CardStatus(None, count)
    if best.missing_required:
        return CardStatus(False, t("card.profiles.none_fits", lang).format(count=count))
    return CardStatus(True, t("card.profiles.fits", lang).format(count=count, profile=quote(best.profile, lang)))


def excel_status(path: Path, data_sheets: list[str], error_sheet: str, lang: str) -> CardStatus:
    """`data_sheets`: the sheet for all rows, or one per profile."""
    if not path.exists():
        return CardStatus(None, t("card.excel.new", lang).format(file=path.name))
    try:
        workbook = load_workbook(path, read_only=True)
        try:
            rows = sum(_data_rows(workbook, sheet) for sheet in data_sheets)
            problems = _data_rows(workbook, error_sheet)
        finally:
            workbook.close()
    except Exception:  # locked, damaged or not an Excel file: the card is only informational
        return CardStatus(None, path.name)
    text = t("card.excel.rows", lang).format(file=path.name, rows=rows)
    if problems:
        text += t("card.excel.problems", lang).format(count=problems)
    return CardStatus(True, text)


def _data_rows(workbook, sheet_name: str) -> int:
    if sheet_name not in workbook.sheetnames:
        return 0
    return max((workbook[sheet_name].max_row or 1) - 1, 0)
