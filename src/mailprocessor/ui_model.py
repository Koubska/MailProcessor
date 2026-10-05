"""Form handling and status texts for the GUI, without tkinter (testable).

The GUI keeps every input as a plain value in a dict (`FormValues`). `config_from_form` turns it into an
`AppConfig` and reports problems per input, so the GUI can show them next to the field and save
automatically only when everything is valid.
"""

from __future__ import annotations

import re
from collections.abc import Iterator
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from pydantic import ValidationError

from mailprocessor.config import (
    DEFAULT_PROFILE_NAME,
    SHEET_NAME_MAX_LENGTH,
    AppConfig,
    AppSection,
    EmlSourceConfig,
    FieldRule,
    ImapSourceConfig,
    MailFilter,
    ParsingRules,
    Profile,
    SourceConfig,
    profile_name_problem,
)
from mailprocessor.gui import setting_path, split_labels
from mailprocessor.i18n import t
from mailprocessor.parser import ParseResult
from mailprocessor.preview import RulePreview, sample_files
from mailprocessor.processor import Problem

FormValues = dict[str, str | bool]

DEFAULT_IMAP_PORTS = {True: "993", False: "143"}


def form_values(config: AppConfig) -> FormValues:
    app, source = config.app, config.source
    eml = source.eml or EmlSourceConfig(folder="", glob="*.eml")
    imap = source.imap or ImapSourceConfig(host="", username="")
    return {
        "source_type": source.type,
        "eml_folder": eml.folder,
        "eml_glob": eml.glob,
        "imap_host": imap.host,
        "imap_port": str(imap.port),
        "imap_username": imap.username,
        "imap_mailbox": imap.mailbox,
        "imap_use_ssl": imap.use_ssl,
        "filter_subject": "; ".join(config.filter.subject),
        "filter_sender": "; ".join(config.filter.sender),
        "output_xlsx": app.output_xlsx,
        "sheet_data": app.sheet_data,
        "sheet_errors": app.sheet_errors,
        "profile_sheets": app.profile_sheets,
        "log_level": app.log_level,
        "sqlite_path": app.sqlite_path,
        "max_age_days": str(app.max_age_days),
        "max_messages": str(app.max_messages),
        "dry_run": app.dry_run,
    }


def _text(values: FormValues, key: str) -> str:
    return str(values.get(key, "")).strip()


def _count(values: FormValues, key: str, errors: dict[str, str], lang: str) -> int:
    raw = _text(values, key) or "0"
    if not raw.isdigit():
        errors[key] = t("error.form.number", lang)
        return 0
    return int(raw)


def config_from_form(values: FormValues, lang: str) -> tuple[AppConfig | None, dict[str, str]]:
    """The config for the inputs, or None plus a message per invalid input (key = input name)."""
    errors: dict[str, str] = {}
    source_type = _text(values, "source_type") or "eml"
    use_imap = source_type == "imap"

    if source_type == "eml" and not _text(values, "eml_folder"):
        errors["eml_folder"] = t("error.form.folder_empty", lang)
    if use_imap and not _text(values, "imap_host"):
        errors["imap_host"] = t("error.form.host_empty", lang)
    if use_imap and not _text(values, "imap_username"):
        errors["imap_username"] = t("error.form.username_empty", lang)
    port_text = _text(values, "imap_port") or DEFAULT_IMAP_PORTS[bool(values.get("imap_use_ssl", True))]
    port = int(port_text) if port_text.isdigit() and 0 < int(port_text) < 65536 else None
    if port is None and use_imap:
        errors["imap_port"] = t("error.form.port", lang)

    output = _text(values, "output_xlsx")
    if not output:
        errors["output_xlsx"] = t("error.form.output_empty", lang)
    elif not output.lower().endswith(".xlsx"):
        errors["output_xlsx"] = t("error.form.output_suffix", lang)
    sheet_data, sheet_errors = _text(values, "sheet_data"), _text(values, "sheet_errors")
    for key, name in (("sheet_data", sheet_data), ("sheet_errors", sheet_errors)):
        if not name:
            errors[key] = t("error.form.sheet_empty", lang)
    if sheet_data and sheet_data == sheet_errors:
        errors["sheet_errors"] = t("error.form.sheets_equal", lang)
    if not _text(values, "sqlite_path"):
        errors["sqlite_path"] = t("error.form.sqlite_empty", lang)
    max_age_days = _count(values, "max_age_days", errors, lang)
    max_messages = _count(values, "max_messages", errors, lang)
    if errors:
        return None, errors

    # The inactive source's settings are kept, so switching back and forth loses nothing.
    eml = None
    if source_type == "eml" or _text(values, "eml_folder"):
        eml = EmlSourceConfig(folder=_text(values, "eml_folder"), glob=_text(values, "eml_glob") or "*.eml")
    imap = None
    if use_imap or _text(values, "imap_host"):
        imap = ImapSourceConfig(
            host=_text(values, "imap_host"),
            port=port or int(DEFAULT_IMAP_PORTS[True]),
            username=_text(values, "imap_username"),
            mailbox=_text(values, "imap_mailbox") or "INBOX",
            use_ssl=bool(values.get("imap_use_ssl", True)),
        )
    try:
        config = AppConfig(
            app=AppSection(
                log_level=_text(values, "log_level") or "INFO",
                sqlite_path=_text(values, "sqlite_path"),
                output_xlsx=output,
                sheet_data=sheet_data,
                sheet_errors=sheet_errors,
                profile_sheets="per_profile" if _text(values, "profile_sheets") == "per_profile" else "shared",
                dry_run=bool(values.get("dry_run", False)),
                max_messages=max_messages,
                max_age_days=max_age_days,
            ),
            source=SourceConfig(type=source_type, eml=eml, imap=imap),
            filter=MailFilter(
                subject=split_labels(_text(values, "filter_subject")),
                sender=split_labels(_text(values, "filter_sender")),
            ),
        )
    except ValidationError as exc:
        return None, {"_": exc.errors()[0]["msg"].removeprefix("Value error, ")}
    return config, {}


@dataclass(frozen=True)
class CardStatus:
    """One status card on the start page. ok: True = ready, False = needs attention, None = neutral."""

    ok: bool | None
    text: str


def _quote(text: str, lang: str) -> str:
    return f"„{text}“" if lang == "de" else f"“{text}”"


def mail_count_in_folder(values: FormValues, config_dir: Path) -> int | None:
    """Number of mail files in the configured folder, or None if the folder does not exist."""
    folder = setting_path(_text(values, "eml_folder") or ".", config_dir)
    if not folder.is_dir():
        return None
    return len(sample_files(folder, _text(values, "eml_glob") or "*.eml"))


def mails_status(values: FormValues, config_dir: Path, password_given: bool, lang: str) -> CardStatus:
    status = _source_status(values, config_dir, password_given, lang)
    if not status.ok:
        return status
    # Ready: say which mails count, so nobody wonders why some are not in the workbook.
    or_word = f" {t('card.mails.filter_or', lang)} "
    notes = [
        t(key, lang).format(entries=or_word.join(_quote(entry, lang) for entry in entries))
        for key, entries in (
            ("card.mails.filter_subject", split_labels(_text(values, "filter_subject"))),
            ("card.mails.filter_sender", split_labels(_text(values, "filter_sender"))),
        )
        if entries
    ]
    return CardStatus(status.ok, " · ".join([status.text, *notes]))


def _source_status(values: FormValues, config_dir: Path, password_given: bool, lang: str) -> CardStatus:
    if _text(values, "source_type") == "imap":
        host, user = _text(values, "imap_host"), _text(values, "imap_username")
        if not host or not user:
            return CardStatus(False, t("card.mails.imap_incomplete", lang))
        if not password_given:
            return CardStatus(False, t("card.mails.imap_password", lang).format(user=user))
        return CardStatus(True, t("card.mails.imap_ready", lang).format(user=user, host=host))
    folder_name = _quote(Path(_text(values, "eml_folder") or ".").name or _text(values, "eml_folder"), lang)
    count = mail_count_in_folder(values, config_dir)
    if count is None:
        return CardStatus(False, t("card.mails.folder_missing", lang).format(folder=folder_name))
    if count == 0:
        return CardStatus(False, t("card.mails.folder_empty", lang).format(folder=folder_name))
    return CardStatus(True, t("card.mails.folder_ready", lang).format(folder=folder_name, count=count))


def fields_status(rules: list[FieldRule], previews: list[RulePreview] | None, lang: str) -> CardStatus:
    if not rules:
        return CardStatus(False, t("card.fields.none", lang))
    count = t("card.fields.count", lang).format(count=len(rules))
    if previews is None:
        return CardStatus(None, count)
    missing = [
        rule.column for rule, preview in zip(rules, previews, strict=True) if rule.required and not preview.found
    ]
    if missing:
        columns = ", ".join(_quote(column, lang) for column in missing)
        return CardStatus(False, t("card.fields.missing", lang).format(count=count, columns=columns))
    return CardStatus(True, t("card.fields.ok", lang).format(count=count))


def profiles_status(profile_count: int, field_count: int, best: ParseResult | None, lang: str) -> CardStatus:
    """The fields card when there are several profiles: which profile the sample mail fits."""
    count = t("card.profiles.count", lang).format(profiles=profile_count, fields=field_count)
    if best is None:
        return CardStatus(None, count)
    if best.missing_required:
        return CardStatus(False, t("card.profiles.none_fits", lang).format(count=count))
    return CardStatus(True, t("card.profiles.fits", lang).format(count=count, profile=_quote(best.profile, lang)))


def excel_status(path: Path, data_sheets: list[str], error_sheet: str, lang: str) -> CardStatus:
    """`data_sheets`: the sheet for all rows, or one per profile."""
    if not path.exists():
        return CardStatus(None, t("card.excel.new", lang).format(file=path.name))
    try:
        workbook = load_workbook(path, read_only=True)
        try:
            rows = sum(_data_rows(workbook, sheet) for sheet in data_sheets)
            problems = _data_rows(workbook, error_sheet)
        finally:
            workbook.close()
    except Exception:  # locked, damaged or not an Excel file: the card is only informational
        return CardStatus(None, path.name)
    text = t("card.excel.rows", lang).format(file=path.name, rows=rows)
    if problems:
        text += t("card.excel.problems", lang).format(count=problems)
    return CardStatus(True, text)


def _data_rows(workbook, sheet_name: str) -> int:
    if sheet_name not in workbook.sheetnames:
        return 0
    return max((workbook[sheet_name].max_row or 1) - 1, 0)


def problem_text(problem: Problem, lang: str) -> str:
    """What is wrong with a failed mail, for the problem list."""
    if problem.missing:
        columns = ", ".join(_quote(column, lang) for column in problem.missing)
        if problem.profile:
            return t("problems.missing_profile", lang).format(profile=_quote(problem.profile, lang), columns=columns)
        return t("problems.missing", lang).format(columns=columns)
    if problem.body is None:
        return t("problems.unreadable_short", lang)
    return problem.reason


class RuleList:
    """The fields edited in the GUI. Every change is validated like the rules file; a removal can be undone."""

    def __init__(self, rules: list[FieldRule]) -> None:
        self.rules = list(rules)
        self._removed: tuple[int, FieldRule] | None = None

    def __len__(self) -> int:
        return len(self.rules)

    def __iter__(self) -> Iterator[FieldRule]:
        return iter(self.rules)

    def __getitem__(self, index: int) -> FieldRule:
        return self.rules[index]

    def columns(self, except_index: int | None = None) -> list[str]:
        return [rule.column for index, rule in enumerate(self.rules) if index != except_index]

    def parsing_rules(self) -> ParsingRules:
        return ParsingRules(fields=self.rules)

    def _commit(self, rules: list[FieldRule]) -> None:
        if rules:
            ParsingRules(fields=rules)  # raises ValueError, e.g. for duplicates or the reserved column name
        self.rules = rules
        self._removed = None

    def add(self, rule: FieldRule) -> int:
        self._commit([*self.rules, rule])
        return len(self.rules) - 1

    def replace(self, index: int, rule: FieldRule) -> bool:
        """Put `rule` at `index`; returns False (and changes nothing) if it is the same rule."""
        if self.rules[index] == rule:
            return False
        self._commit([*self.rules[:index], rule, *self.rules[index + 1 :]])
        return True

    def move(self, index: int, offset: int) -> int | None:
        new_index = index + offset
        if not 0 <= new_index < len(self.rules):
            return None
        rules = list(self.rules)
        rules[index], rules[new_index] = rules[new_index], rules[index]
        self._commit(rules)
        return new_index

    def remove(self, index: int) -> FieldRule:
        removed = self.rules[index]
        self.rules = [*self.rules[:index], *self.rules[index + 1 :]]
        self._removed = (index, removed)
        return removed

    @property
    def can_undo(self) -> bool:
        return self._removed is not None

    def undo_remove(self) -> int | None:
        """Put the last removed field back where it was; None if nothing to undo or it would be invalid now."""
        if self._removed is None:
            return None
        index, rule = self._removed
        index = min(index, len(self.rules))
        try:
            self._commit([*self.rules[:index], rule, *self.rules[index:]])
        except ValueError:
            self._removed = None
            return None
        return index


class ProfileList:
    """The profiles edited in the GUI, each with its own `RuleList`. The Felder tab shows one at a time.

    A profile without fields is kept while editing but left out when saving or running.
    """

    def __init__(self, profiles: list[Profile]) -> None:
        self.names = [profile.name for profile in profiles] or [DEFAULT_PROFILE_NAME]
        self.fields = [RuleList(profile.fields) for profile in profiles] or [RuleList([])]
        self.index = 0

    def __len__(self) -> int:
        return len(self.names)

    @property
    def current(self) -> RuleList:
        return self.fields[self.index]

    @property
    def current_name(self) -> str:
        return self.names[self.index]

    @property
    def field_count(self) -> int:
        return sum(len(rules) for rules in self.fields)

    def select(self, index: int) -> None:
        if 0 <= index < len(self.names):
            self.index = index

    def find(self, name: str) -> int | None:
        return self.names.index(name) if name in self.names else None

    def name_problem(self, name: str, lang: str, except_index: int | None = None) -> str | None:
        """Why `name` cannot be used for a new or renamed profile, in the user's language; None if it can."""
        name = name.strip()
        if not name:
            return t("error.profile.name_empty", lang)
        if len(name) > SHEET_NAME_MAX_LENGTH:
            return t("error.profile.name_long", lang).format(max=SHEET_NAME_MAX_LENGTH)
        if profile_name_problem(name):
            return t("error.profile.name_chars", lang)
        taken = [other for index, other in enumerate(self.names) if index != except_index]
        if name.casefold() in (other.casefold() for other in taken):
            return t("error.profile.name_taken", lang).format(name=_quote(name, lang))
        return None

    def add(self, name: str, lang: str) -> int:
        """Add an empty profile and show it; raises ValueError with a message for the user."""
        problem = self.name_problem(name, lang)
        if problem:
            raise ValueError(problem)
        self.names.append(name.strip())
        self.fields.append(RuleList([]))
        self.index = len(self.names) - 1
        return self.index

    def rename(self, name: str, lang: str) -> None:
        problem = self.name_problem(name, lang, except_index=self.index)
        if problem:
            raise ValueError(problem)
        self.names[self.index] = name.strip()

    def remove(self) -> str:
        """Remove the shown profile (never the last one) and show its neighbour."""
        if len(self.names) == 1:
            raise ValueError("the last profile cannot be removed")
        name = self.names.pop(self.index)
        self.fields.pop(self.index)
        self.index = min(self.index, len(self.names) - 1)
        return name

    def other_columns(self) -> list[str]:
        """Columns of the other profiles, each once, in profile order."""
        columns = (
            rule.column for index, rules in enumerate(self.fields) if index != self.index for rule in rules
        )
        return list(dict.fromkeys(columns))

    def column_suggestions(self, typed: str, except_index: int | None = None) -> list[str]:
        """Column names of other profiles to offer in the shown profile: not used there yet, containing `typed`.

        `except_index` is the field being edited; its own name stays available.
        """
        own = set(self.current.columns(except_index=except_index))
        wanted = typed.strip().casefold()
        return [column for column in self.other_columns() if column not in own and wanted in column.casefold()]

    def profiles_with_column(self, column: str) -> list[str]:
        """The other profiles that have exactly this column."""
        return [
            name
            for index, (name, rules) in enumerate(zip(self.names, self.fields, strict=True))
            if index != self.index and column in rules.columns()
        ]

    def rule_for_column(self, column: str) -> FieldRule | None:
        """How another profile finds this column (the first that has it), to start from in the shown profile."""
        for index, rules in enumerate(self.fields):
            if index != self.index:
                for rule in rules:
                    if rule.column == column:
                        return rule
        return None

    def profiles(self) -> list[Profile]:
        """The profiles to save: all that have fields."""
        return [
            Profile(name=name, fields=rules.rules) for name, rules in zip(self.names, self.fields, strict=True) if rules
        ]

    def parsing_rules(self) -> ParsingRules:
        """Raises ValueError if no profile has fields yet."""
        return ParsingRules(profiles=self.profiles())


def complete_column(typed: str, suggestions: list[str]) -> str | None:
    """The first suggestion that continues what was typed (ignoring case), for inline completion."""
    wanted = typed.casefold()
    if not wanted:
        return None
    for suggestion in suggestions:
        if suggestion.casefold().startswith(wanted) and len(suggestion) > len(typed):
            return suggestion
    return None


def _column_key(column: str) -> str:
    """Spelling-insensitive form of a column name, to spot near-duplicates such as "Telefon-Nr" / "telefon nr"."""
    return re.sub(r"[\s\-_.:/]+", "", column.casefold())


@dataclass(frozen=True)
class ColumnNote:
    text: str
    warning: bool = False


def column_note(profiles: ProfileList, column: str, per_profile_sheets: bool, lang: str) -> ColumnNote:
    """Where the column ends up in Excel, shown under the column name when there are several profiles."""
    column = column.strip()
    if len(profiles) < 2 or not column:
        return ColumnNote("")
    if per_profile_sheets:
        return ColumnNote(t("column.note.per_profile", lang).format(sheet=_quote(profiles.current_name, lang)))
    same = profiles.profiles_with_column(column)
    if same:
        key = "column.note.shared_one" if len(same) == 1 else "column.note.shared_many"
        return ColumnNote(t(key, lang).format(profiles=", ".join(_quote(name, lang) for name in same)))
    similar = [other for other in profiles.other_columns() if _column_key(other) == _column_key(column)]
    if similar:
        return ColumnNote(
            t("column.note.similar", lang).format(
                column=_quote(similar[0], lang),
                profiles=", ".join(_quote(name, lang) for name in profiles.profiles_with_column(similar[0])),
            ),
            warning=True,
        )
    return ColumnNote(t("column.note.own", lang))


@dataclass(frozen=True)
class Shortcut:
    sequences: tuple[str, ...]  # tkinter event sequences
    label: str  # shown in hints, e.g. "Strg+Enter"


def shortcuts(platform: str, lang: str) -> dict[str, Shortcut]:
    """Keyboard shortcuts for the main actions; Command on macOS, Control elsewhere."""
    if platform == "darwin":
        return {
            "run": Shortcut(("<Command-Return>",), "⌘↩"),
            "test_run": Shortcut(("<Command-Shift-Return>",), "⇧⌘↩"),
            "open_excel": Shortcut(("<Command-e>", "<Command-E>"), "⌘E"),
        }
    control, shift = ("Strg", "Umschalt") if lang == "de" else ("Ctrl", "Shift")
    return {
        "run": Shortcut(("<Control-Return>",), f"{control}+Enter"),
        "test_run": Shortcut(("<Control-Shift-Return>",), f"{control}+{shift}+Enter"),
        "open_excel": Shortcut(("<Control-e>", "<Control-E>"), f"{control}+E"),
    }
