"""Columns, times and helpers shared by the workbook tests."""

from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from mailprocessor.excel_writer import (
    CONTENT_COLUMN,
    RECEIVED_COLUMN,
    TRANSFERRED_COLUMN,
    ErrorEntry,
    ExcelOutput,
)

DATA_COLUMNS = ["Mail-Adresse", "Name", "Kurs", "Zeit", "Telefonnummer"]


FIXED_COLUMNS = [RECEIVED_COLUMN, TRANSFERRED_COLUMN, CONTENT_COLUMN]


RUN_AT = datetime(2026, 10, 6, 9, 30)


RECEIVED_AT = datetime(2026, 10, 5, 14, 0)


def _output(path: Path, columns: list[str]) -> ExcelOutput:
    return ExcelOutput(path, {"daten": columns}, "fehler", now=RUN_AT)


def _rows(path: Path, sheet: str = "daten") -> list[tuple]:
    return list(load_workbook(path)[sheet].iter_rows(values_only=True))


def _error(key: str, reason: str = "Pflichtfelder nicht gefunden", missing: tuple[str, ...] = ("Telefonnummer",)):
    return ErrorEntry(
        key=key,
        name=f"{key}.eml",
        sender="Max Mustermann <max@example.com>",
        subject="Anmeldung",
        received=RECEIVED_AT,
        missing=missing,
        reason=reason,
    )
