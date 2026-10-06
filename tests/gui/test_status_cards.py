"""The status cards on the Start tab."""

from pathlib import Path

from openpyxl import Workbook

from mailprocessor.config import FieldRule
from mailprocessor.gui.config_files import default_config as _default_config
from mailprocessor.gui.form import form_values
from mailprocessor.gui.preview import RulePreview
from mailprocessor.gui.status_cards import CardStatus, excel_status, fields_status, mails_status, profiles_status
from mailprocessor.parser import ParseResult


def test_mails_status_for_folder(tmp_path: Path) -> None:
    values = {**form_values(_default_config()), "eml_folder": "./mails"}

    assert mails_status(values, tmp_path, False, "de").ok is False  # missing
    (tmp_path / "mails").mkdir()
    assert "ist leer" in mails_status(values, tmp_path, False, "de").text
    (tmp_path / "mails" / "a.eml").write_text("x", encoding="utf-8")
    status = mails_status(values, tmp_path, False, "de")
    assert status.ok and status.text == "Ordner „mails“ – 1 E-Mail(s)"


def test_mails_status_names_an_active_filter(tmp_path: Path) -> None:
    (tmp_path / "mails").mkdir()
    (tmp_path / "mails" / "a.eml").write_text("x", encoding="utf-8")
    values = {**form_values(_default_config()), "eml_folder": "./mails", "filter_subject": "Kontakt; Anmeldung"}

    assert mails_status(values, tmp_path, False, "de").text == (
        "Ordner „mails“ – 1 E-Mail(s) · nur mit Betreff „Kontakt“ oder „Anmeldung“"
    )
    values = {**values, "source_type": "imap", "filter_subject": "", "filter_sender": "schule@example.com"}
    assert mails_status(values, tmp_path, True, "en").text == (
        "Mailbox user@example.com on imap.example.com · only from “schule@example.com”"
    )


def test_mails_status_for_imap_asks_for_the_password() -> None:
    values = {**form_values(_default_config()), "source_type": "imap"}

    assert "Passwort" in mails_status(values, Path("."), False, "de").text
    assert mails_status(values, Path("."), True, "de").text == "Postfach user@example.com auf imap.example.com"


def test_fields_status() -> None:
    rules = [FieldRule(column="Name", type="label", label="Name"), FieldRule(column="Tel", type="label", label="Tel")]

    assert fields_status([], None, "de").ok is False
    assert fields_status(rules, None, "de").ok is None
    assert fields_status(rules, [RulePreview("a"), RulePreview("b")], "de").ok is True
    missing = fields_status(rules, [RulePreview("a"), RulePreview(None)], "de")
    assert missing.ok is False and missing.text == "2 Feld(er) – in der Beispiel-Mail fehlt „Tel“"


def test_excel_status_counts_rows(tmp_path: Path) -> None:
    path = tmp_path / "out.xlsx"
    assert "wird beim ersten Übertragen angelegt" in excel_status(path, ["daten"], "fehler", "de").text

    workbook = Workbook()
    workbook.active.title = "daten"
    workbook.active.append(["Name"])
    workbook.active.append(["A"])
    workbook.active.append(["B"])
    workbook.create_sheet("fehler").append(["reason"])
    workbook["fehler"].append(["x"])
    workbook.save(path)

    assert excel_status(path, ["daten"], "fehler", "de").text == "out.xlsx – 2 Zeile(n), 1 E-Mail(s) mit Problemen"
    (tmp_path / "broken.xlsx").write_text("not excel", encoding="utf-8")
    assert excel_status(tmp_path / "broken.xlsx", ["daten"], "fehler", "de").ok is None


def test_excel_status_adds_up_the_profile_sheets(tmp_path: Path) -> None:
    path = tmp_path / "out.xlsx"
    workbook = Workbook()
    workbook.active.title = "Anmeldung"
    workbook.active.append(["Name"])
    workbook.active.append(["A"])
    abmeldung = workbook.create_sheet("Abmeldung")
    abmeldung.append(["Name"])
    abmeldung.append(["B"])
    abmeldung.append(["C"])
    workbook.save(path)

    assert excel_status(path, ["Anmeldung", "Abmeldung"], "fehler", "de").text == "out.xlsx – 3 Zeile(n)"


def test_profiles_status() -> None:
    fits = ParseResult({"Kurs": "Judo"}, [], None, "Anmeldung")
    misses = ParseResult({}, ["Kurs"], "missing", "Anmeldung")

    assert profiles_status(2, 7, None, "de") == CardStatus(None, "2 Profile, 7 Feld(er)")
    assert profiles_status(2, 7, fits, "de") == CardStatus(
        True, "2 Profile, 7 Feld(er) – die Beispiel-Mail passt zu „Anmeldung“"
    )
    assert profiles_status(2, 7, misses, "de").ok is False
