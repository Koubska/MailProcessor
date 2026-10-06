"""Config, rules and test mails shared by the pipeline tests."""

from pathlib import Path

from openpyxl import load_workbook

from mailprocessor.config import (
    AppConfig,
    AppSection,
    EmlSourceConfig,
    FieldRule,
    ParsingRules,
    SourceConfig,
)


def _build_config(tmp_path: Path, max_messages: int = 0) -> AppConfig:
    return AppConfig(
        app=AppSection(
            log_level="INFO",
            sqlite_path=str(tmp_path / "data" / "ledger.db"),
            output_xlsx=str(tmp_path / "out" / "mail_export.xlsx"),
            sheet_data="daten",
            sheet_errors="fehler",
            dry_run=False,
            max_messages=max_messages,
            max_age_days=0,
        ),
        source=SourceConfig(
            type="eml",
            eml=EmlSourceConfig(folder=str(tmp_path / "inbox"), glob="*.eml"),
        ),
    )


def _build_rules() -> ParsingRules:
    return ParsingRules(
        fields=[
            FieldRule(column="Mail-Adresse", pattern=r"(?im)^\s*Von:\s*.*?<([^>]+)>\s*$", required=True),
            FieldRule(column="Name", pattern=r"(?im)mein\s+Kind\s+(.+?)\s+f[uü]r\s+folgendes\s+Angebot", required=True),
            FieldRule(column="Kurs", pattern=r"(?im)^\s*Angebot:\s*(.+?)\s*$", required=True),
            FieldRule(column="Zeit", pattern=r"(?im)^\s*Tag:\s*(.+?)\s*$", required=True),
            FieldRule(column="Telefonnummer", pattern=r"(?im)^\s*Telefonnummer:\s*(.+?)\s*$", required=True),
        ]
    )


def _write_eml(path: Path, body_lines: list[str], message_id: str, subject: str = "Schnuppernachmittag") -> None:
    path.write_text(
        "\n".join(
            [
                f"Message-ID: <{message_id}>",
                "From: Max Mustermann <max.mustermann@mail.com>",
                f"Subject: {subject}",
                "Date: Mon, 23 Nov 2026 14:00:00 +0100",
                "Content-Type: text/plain; charset=utf-8",
                "",
                *body_lines,
            ]
        ),
        encoding="utf-8",
    )


def _write_eml_with_date(path: Path, body_lines: list[str], message_id: str, date_header: str) -> None:
    path.write_text(
        "\n".join(
            [
                f"Message-ID: <{message_id}>",
                "From: Max Mustermann <max.mustermann@mail.com>",
                "Subject: Schnuppernachmittag",
                f"Date: {date_header}",
                "Content-Type: text/plain; charset=utf-8",
                "",
                *body_lines,
            ]
        ),
        encoding="utf-8",
    )


GOOD_BODY = [
    "Von: Max Mustermann <max.mustermann@mail.com>",
    "hiermit möchte ich mein Kind Jan Must+ für folgendes Angebot anmelden:",
    "Tag: Mo, 23.11.2026 (14:15-16:00 Uhr)",
    "Angebot: Experimente",
    "Telefonnummer: 1234 567890",
]


def _snapshot(folder: Path) -> dict[str, tuple[bytes, int, int]]:
    return {
        path.name: (path.read_bytes(), path.stat().st_mtime_ns, path.stat().st_mode)
        for path in sorted(folder.iterdir())
    }


def _inbox_with_good_mails(tmp_path: Path, count: int) -> Path:
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    for index in range(count):
        _write_eml(inbox / f"mail-{index}.eml", GOOD_BODY, f"mail-{index}@example.com")
    return inbox


def _data_rows(tmp_path: Path) -> list[dict]:
    sheet = load_workbook(tmp_path / "out" / "mail_export.xlsx")["daten"]
    header, *rows = sheet.iter_rows(values_only=True)
    return [dict(zip(header, row, strict=True)) for row in rows]


def _error_rows(tmp_path: Path) -> list[tuple]:
    return list(load_workbook(tmp_path / "out" / "mail_export.xlsx")["fehler"].iter_rows(min_row=2, values_only=True))
