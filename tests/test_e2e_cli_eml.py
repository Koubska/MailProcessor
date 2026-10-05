from pathlib import Path

from openpyxl import load_workbook
from typer.testing import CliRunner

import mailprocessor.main as main_module


def _write_eml(path: Path, message_id: str, body_lines: list[str]) -> None:
    path.write_text(
        "\n".join(
            [
                f"Message-ID: <{message_id}>",
                "From: Max Mustermann <max.mustermann@mail.com>",
                "Subject: Schnuppernachmittag",
                "Date: Mon, 23 Nov 2026 14:00:00 +0100",
                "Content-Type: text/plain; charset=utf-8",
                "",
                *body_lines,
            ]
        ),
        encoding="utf-8",
    )


def _write_config(path: Path, *, db_path: Path, output_path: Path, inbox: Path) -> None:
    path.write_text(
        "\n".join(
            [
                "[app]",
                'log_level = "INFO"',
                f'sqlite_path = "{db_path.as_posix()}"',
                f'output_xlsx = "{output_path.as_posix()}"',
                'sheet_data = "daten"',
                'sheet_errors = "fehler"',
                "dry_run = false",
                "max_messages = 0",
                "",
                "[source]",
                'type = "eml"',
                "",
                "[source.eml]",
                f'folder = "{inbox.as_posix()}"',
                'glob = "*.eml"',
            ]
        ),
        encoding="utf-8",
    )


def _write_rules(path: Path) -> None:
    path.write_text(
        "\n".join(
            [
                "[[fields]]",
                'column = "Mail-Adresse"',
                r"pattern = '(?im)^\s*Von:\s*.*?<([^>]+)>\s*$'",
                "required = true",
                "",
                "[[fields]]",
                'column = "Name"',
                r"pattern = '(?im)mein\s+Kind\s+(.+?)\s+f[uü]r\s+folgendes\s+Angebot'",
                "required = true",
                "",
                "[[fields]]",
                'column = "Kurs"',
                r"pattern = '(?im)^\s*Angebot:\s*(.+?)\s*$'",
                "required = true",
                "",
                "[[fields]]",
                'column = "Zeit"',
                r"pattern = '(?im)^\s*Tag:\s*(.+?)\s*$'",
                "required = true",
                "",
                "[[fields]]",
                'column = "Telefonnummer"',
                r"pattern = '(?im)^\s*Telefonnummer:\s*(.+?)\s*$'",
                "required = true",
            ]
        ),
        encoding="utf-8",
    )


def test_cli_eml_end_to_end_with_excel_and_idempotency(tmp_path: Path) -> None:
    inbox = tmp_path / "inbox"
    inbox.mkdir()

    _write_eml(
        inbox / "good.eml",
        "good@example.com",
        [
            "Von: Max Mustermann <max.mustermann@mail.com>",
            "hiermit möchte ich mein Kind Jan Must+ für folgendes Angebot anmelden:",
            "Tag: Mo, 23.11.2026 (14:15-16:00 Uhr)",
            "Angebot: Experimente",
            "Telefonnummer: 1234 567890",
        ],
    )
    _write_eml(
        inbox / "bad.eml",
        "bad@example.com",
        [
            "Von: Max Mustermann <max.mustermann@mail.com>",
            "hiermit möchte ich mein Kind Jan Must+ für folgendes Angebot anmelden:",
            "Tag: Mo, 23.11.2026 (14:15-16:00 Uhr)",
            "Angebot: Experimente",
        ],
    )

    output_path = tmp_path / "out" / "mail_export.xlsx"
    db_path = tmp_path / "data" / "ledger.db"
    config_path = tmp_path / "config.toml"
    rules_path = tmp_path / "rules.toml"

    _write_config(config_path, db_path=db_path, output_path=output_path, inbox=inbox)
    _write_rules(rules_path)

    runner = CliRunner()
    first_run = runner.invoke(main_module.app, ["--config", str(config_path), "--rules", str(rules_path)])
    assert first_run.exit_code == 0
    assert "seen=2 processed=1 skipped=0 failed=1" in first_run.stdout

    workbook = load_workbook(output_path)
    assert workbook["daten"].max_row == 2
    assert workbook["fehler"].max_row == 2

    # The log file next to the config names the mails but contains no mail content or extracted values.
    log_text = (tmp_path / "logs" / "mailprocessor.log").read_text(encoding="utf-8")
    assert "good.eml" in log_text and "bad.eml" in log_text
    for confidential in ("Jan Must+", "1234 567890", "Experimente", "max.mustermann@mail.com"):
        assert confidential not in log_text

    second_run = runner.invoke(main_module.app, ["--config", str(config_path), "--rules", str(rules_path)])
    assert second_run.exit_code == 0
    assert "seen=2 processed=0 skipped=1 failed=1" in second_run.stdout

    workbook = load_workbook(output_path)
    assert workbook["daten"].max_row == 2
    # The retried failure replaces its error row instead of adding a duplicate.
    assert workbook["fehler"].max_row == 2


def test_cli_eml_many_messages_does_not_duplicate_processed_rows(tmp_path: Path) -> None:
    inbox = tmp_path / "inbox"
    inbox.mkdir()

    total_messages = 120
    for index in range(total_messages):
        _write_eml(
            inbox / f"mail-{index:03d}.eml",
            f"bulk-{index}@example.com",
            [
                f"Von: Schule {index} <schule{index}@example.com>",
                f"hiermit möchte ich mein Kind Kind{index} für folgendes Angebot anmelden:",
                f"Tag: Mo, 23.11.2026 ({14 + (index % 4)}:15-16:00 Uhr)",
                f"Angebot: Kurs {index % 7}",
                f"Telefonnummer: 555 {index:06d}",
            ],
        )

    output_path = tmp_path / "out" / "mail_export.xlsx"
    db_path = tmp_path / "data" / "ledger.db"
    config_path = tmp_path / "config.toml"
    rules_path = tmp_path / "rules.toml"
    _write_config(config_path, db_path=db_path, output_path=output_path, inbox=inbox)
    _write_rules(rules_path)

    runner = CliRunner()
    first_run = runner.invoke(main_module.app, ["--config", str(config_path), "--rules", str(rules_path)])
    assert first_run.exit_code == 0
    assert f"seen={total_messages} processed={total_messages} skipped=0 failed=0" in first_run.stdout

    workbook = load_workbook(output_path)
    data_sheet = workbook["daten"]
    assert data_sheet.max_row == total_messages + 1
    assert workbook["fehler"].max_row == 1
    phone_values_first_run = [data_sheet.cell(row=row_index, column=5).value for row_index in range(2, total_messages + 2)]
    assert len(phone_values_first_run) == total_messages
    assert len(set(phone_values_first_run)) == total_messages

    second_run = runner.invoke(main_module.app, ["--config", str(config_path), "--rules", str(rules_path)])
    assert second_run.exit_code == 0
    assert (
        f"seen={total_messages} processed=0 skipped={total_messages} failed=0"
        in second_run.stdout
    )

    workbook = load_workbook(output_path)
    data_sheet = workbook["daten"]
    assert data_sheet.max_row == total_messages + 1
    assert workbook["fehler"].max_row == 1
    phone_values_second_run = [data_sheet.cell(row=row_index, column=5).value for row_index in range(2, total_messages + 2)]
    assert phone_values_second_run == phone_values_first_run


def test_cli_eml_incremental_run_appends_only_new_messages(tmp_path: Path) -> None:
    inbox = tmp_path / "inbox"
    inbox.mkdir()

    initial_messages = 120
    newly_added_messages = 10
    for index in range(initial_messages):
        _write_eml(
            inbox / f"mail-{index:03d}.eml",
            f"bulk-{index}@example.com",
            [
                f"Von: Schule {index} <schule{index}@example.com>",
                f"hiermit möchte ich mein Kind Kind{index} für folgendes Angebot anmelden:",
                f"Tag: Mo, 23.11.2026 ({14 + (index % 4)}:15-16:00 Uhr)",
                f"Angebot: Kurs {index % 7}",
                f"Telefonnummer: 555 {index:06d}",
            ],
        )

    output_path = tmp_path / "out" / "mail_export.xlsx"
    db_path = tmp_path / "data" / "ledger.db"
    config_path = tmp_path / "config.toml"
    rules_path = tmp_path / "rules.toml"
    _write_config(config_path, db_path=db_path, output_path=output_path, inbox=inbox)
    _write_rules(rules_path)

    runner = CliRunner()
    first_run = runner.invoke(main_module.app, ["--config", str(config_path), "--rules", str(rules_path)])
    assert first_run.exit_code == 0
    assert f"seen={initial_messages} processed={initial_messages} skipped=0 failed=0" in first_run.stdout

    for index in range(initial_messages, initial_messages + newly_added_messages):
        _write_eml(
            inbox / f"mail-{index:03d}.eml",
            f"bulk-{index}@example.com",
            [
                f"Von: Schule {index} <schule{index}@example.com>",
                f"hiermit möchte ich mein Kind Kind{index} für folgendes Angebot anmelden:",
                f"Tag: Mo, 23.11.2026 ({14 + (index % 4)}:15-16:00 Uhr)",
                f"Angebot: Kurs {index % 7}",
                f"Telefonnummer: 555 {index:06d}",
            ],
        )

    second_run = runner.invoke(main_module.app, ["--config", str(config_path), "--rules", str(rules_path)])
    total_seen_second_run = initial_messages + newly_added_messages
    assert second_run.exit_code == 0
    assert (
        f"seen={total_seen_second_run} processed={newly_added_messages} "
        f"skipped={initial_messages} failed=0"
    ) in second_run.stdout

    workbook = load_workbook(output_path)
    data_sheet = workbook["daten"]
    expected_total_rows = initial_messages + newly_added_messages + 1
    assert data_sheet.max_row == expected_total_rows
    assert workbook["fehler"].max_row == 1

    phone_values = [data_sheet.cell(row=row_index, column=5).value for row_index in range(2, expected_total_rows + 1)]
    assert len(phone_values) == initial_messages + newly_added_messages
    assert len(set(phone_values)) == initial_messages + newly_added_messages
