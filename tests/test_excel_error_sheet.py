"""The error sheet: one readable row per failing mail."""

from datetime import datetime
from pathlib import Path

from excel_helpers import DATA_COLUMNS, RECEIVED_AT, RUN_AT, _error, _output, _rows
from openpyxl import Workbook, load_workbook

from mailprocessor.excel_writer import (
    CONTENT_COLUMN,
    ERROR_COLUMNS,
    KEY_COLUMN,
    LEGACY_ERROR_COLUMNS,
    error_key,
)


def test_error_rows_are_readable(tmp_path: Path) -> None:
    path = tmp_path / "output.xlsx"
    output = _output(path, DATA_COLUMNS)

    output.upsert_error(_error("eml|/mails|<a@example.com>", missing=("Telefonnummer", "Kurs")))
    output.save()

    sheet = load_workbook(path)["fehler"]
    row = dict(zip(ERROR_COLUMNS, next(sheet.iter_rows(min_row=2, values_only=True)), strict=True))
    assert row == {
        "E-Mail": "eml|/mails|<a@example.com>.eml",
        "Absender": "Max Mustermann <max@example.com>",
        "Betreff": "Anmeldung",
        "Eingegangen am": RECEIVED_AT,
        "Fehlende Felder": "Telefonnummer, Kurs",
        "Grund": "Pflichtfelder nicht gefunden",
        "Geprüft am": RUN_AT,
        KEY_COLUMN: "eml|/mails|<a@example.com>",
    }
    key_letter = sheet.cell(row=1, column=ERROR_COLUMNS.index(KEY_COLUMN) + 1).column_letter
    assert sheet.column_dimensions[key_letter].hidden  # internal; only used to find the row again


def test_upsert_error_keeps_one_row_per_message(tmp_path: Path) -> None:
    path = tmp_path / "output.xlsx"
    output = _output(path, DATA_COLUMNS)

    output.upsert_error(_error("msg-1", "erster Grund"))
    output.upsert_error(_error("msg-2"))
    output.upsert_error(_error("msg-1", "zweiter Grund"))
    output.save()

    reason, key = ERROR_COLUMNS.index("Grund"), ERROR_COLUMNS.index(KEY_COLUMN)
    rows = _rows(path, "fehler")[1:]
    assert [(row[key], row[reason]) for row in rows] == [
        ("msg-2", "Pflichtfelder nicht gefunden"),
        ("msg-1", "zweiter Grund"),
    ]


def test_remove_errors_for_drops_the_row_of_a_mail_that_now_works(tmp_path: Path) -> None:
    path = tmp_path / "output.xlsx"
    output = _output(path, DATA_COLUMNS)
    output.upsert_error(_error("msg-1"))

    output.remove_errors_for("msg-1")
    output.save()

    assert load_workbook(path)["fehler"].max_row == 1


def test_old_english_error_sheet_is_converted(tmp_path: Path) -> None:
    path = tmp_path / "output.xlsx"
    old = Workbook()
    old.active.title = "daten"
    old.active.append(["Name", CONTENT_COLUMN])
    errors = old.create_sheet("fehler")
    errors.append(LEGACY_ERROR_COLUMNS)
    checked = "2026-10-03T20:00:00+00:00"
    missing = "Required fields missing: Telefonnummer"
    errors.append(["eml", "/mails", "<a@example.com>", "Telefonnummer", missing, checked])
    errors.append(["eml", "/mails", "file:kaputt.eml", "", "Could not read file (OSError): disk", checked])
    old.save(path)

    _output(path, ["Name"]).save()

    header, *data = _rows(path, "fehler")
    rows = [dict(zip(ERROR_COLUMNS, row, strict=True)) for row in data]
    assert list(header) == list(ERROR_COLUMNS)
    assert rows[0]["E-Mail"] == "<a@example.com>"
    assert rows[0]["Fehlende Felder"] == "Telefonnummer"
    assert rows[0]["Grund"] == "Pflichtfelder nicht gefunden"
    assert rows[0][KEY_COLUMN] == error_key("eml", "/mails", "<a@example.com>")
    assert rows[1]["E-Mail"] == "kaputt.eml"
    assert rows[1]["Grund"] == "Could not read file (OSError): disk"
    assert isinstance(rows[0]["Geprüft am"], datetime)

    # The converted row is found again: the next failure of that mail replaces it.
    again = _output(path, ["Name"])
    again.upsert_error(_error(error_key("eml", "/elsewhere", "<a@example.com>")))
    again.save()
    assert len(_rows(path, "fehler")) == 3


def test_error_rows_of_earlier_versions_are_found_without_the_location(tmp_path: Path) -> None:
    # Earlier versions wrote "type|location|identity"; such a row must still be replaced, not duplicated.
    path = tmp_path / "output.xlsx"
    first = _output(path, DATA_COLUMNS)
    first.upsert_error(_error("eml|/old/mails|<a@example.com>"))
    first.upsert_error(_error("eml|/old/mails|<b@example.com>"))
    first.save()

    second = _output(path, DATA_COLUMNS)
    second.upsert_error(_error(error_key("eml", "/new/mails", "<a@example.com>")))
    second.remove_errors_for(error_key("imap", "imap://h:993/INBOX", "<b@example.com>"))  # other source type
    second.save()

    keys = sorted(row[-1] for row in _rows(path, "fehler")[1:])
    assert keys == ["eml|/old/mails|<b@example.com>", "eml|<a@example.com>"]
