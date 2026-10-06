"""Header style and column widths written on save."""

import re
import zipfile
from pathlib import Path

from excel_helpers import DATA_COLUMNS, RECEIVED_AT, _output
from openpyxl import load_workbook
from openpyxl.worksheet.dimensions import ColumnDimension

from mailprocessor.excel_writer import (
    CONTENT_COLUMN_WIDTH,
    DATE_COLUMN_WIDTH,
    MAX_COLUMN_WIDTH,
    MIN_COLUMN_WIDTH,
)


def test_save_formats_header_and_column_widths(tmp_path: Path) -> None:
    path = tmp_path / "output.xlsx"
    output = _output(path, DATA_COLUMNS)
    output.append_data(
        {"Mail-Adresse": "max.mustermann@mail.com", "Name": "x" * 200}, "Text\n" * 50, received=RECEIVED_AT
    )
    output.save()

    sheet = load_workbook(path)["daten"]
    assert all(cell.font.bold for cell in sheet[1])
    widths = {letter: sheet.column_dimensions[letter].width for letter in "ABCFGH"}
    assert widths["A"] == len("max.mustermann@mail.com") + 2
    assert widths["B"] == MAX_COLUMN_WIDTH
    assert widths["C"] == MIN_COLUMN_WIDTH  # "Kurs", empty
    assert widths["F"] == widths["G"] == DATE_COLUMN_WIDTH
    assert widths["H"] == CONTENT_COLUMN_WIDTH
    assert not sheet["H2"].alignment.wrap_text  # one row per mail, however long the text


def test_save_keeps_column_widths_set_by_the_user(tmp_path: Path) -> None:
    path = tmp_path / "output.xlsx"
    _output(path, DATA_COLUMNS).save()
    workbook = load_workbook(path)
    workbook["daten"].column_dimensions["B"].width = 33
    workbook.save(path)

    output = _output(path, DATA_COLUMNS)
    output.append_data({"Name": "x" * 200}, "Text")
    output.save()

    assert load_workbook(path)["daten"].column_dimensions["B"].width == 33


def _column_ranges(path: Path, sheet_index: int = 1) -> list[tuple[int, int]]:
    xml = zipfile.ZipFile(path).read(f"xl/worksheets/sheet{sheet_index}.xml").decode()
    return [(int(low), int(high)) for low, high in re.findall(r'<col [^>]*min="(\d+)" max="(\d+)"', xml)]


def test_save_keeps_a_width_set_for_several_columns_without_overlapping_ranges(tmp_path: Path) -> None:
    # Excel stores one width for several selected columns as one range (<col min="2" max="4">).
    # Widths added on top for columns inside it make Excel report the file as damaged.
    path = tmp_path / "output.xlsx"
    _output(path, DATA_COLUMNS).save()
    workbook = load_workbook(path)
    sheet = workbook["daten"]
    for letter in list(sheet.column_dimensions):
        del sheet.column_dimensions[letter]
    sheet.column_dimensions["B"] = ColumnDimension(sheet, index="B", min=2, max=4, width=33, customWidth=True)
    workbook.save(path)

    output = _output(path, DATA_COLUMNS)
    output.append_data({"Name": "x" * 200, "Kurs": "y" * 200}, "Text")
    output.save()

    ranges = sorted(_column_ranges(path))
    covered = [column for low, high in ranges for column in range(low, high + 1)]
    assert len(covered) == len(set(covered)), ranges
    assert (2, 4) in ranges
    assert load_workbook(path)["daten"].column_dimensions["B"].width == 33
