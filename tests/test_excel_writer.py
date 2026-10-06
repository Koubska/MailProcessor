"""The data sheets: columns found by name, user columns kept, values stored safely."""

from pathlib import Path

import pytest
from excel_helpers import DATA_COLUMNS, FIXED_COLUMNS, RECEIVED_AT, RUN_AT, _error, _output, _rows
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Border, Side

from mailprocessor.errors import SheetHeaderError, WorkbookUnreadableError
from mailprocessor.excel_writer import (
    DATE_FORMAT,
    ERROR_COLUMNS,
    EXCEL_CELL_LIMIT,
    ExcelOutput,
)


def test_several_data_sheets_get_their_own_columns_and_rows(tmp_path: Path) -> None:
    path = tmp_path / "output.xlsx"
    output = ExcelOutput(path, {"Anmeldung": ["Name", "Kurs"], "Abmeldung": ["Name", "Grund"]}, "fehler", now=RUN_AT)

    output.append_data({"Name": "Eva", "Grund": "Umzug"}, "text", sheet="Abmeldung")
    output.append_data({"Name": "Max", "Kurs": "Judo"}, "text")  # first sheet by default
    output.save()

    assert load_workbook(path).sheetnames == ["Anmeldung", "Abmeldung", "fehler"]
    assert _rows(path, "Anmeldung")[0] == ("Name", "Kurs", *FIXED_COLUMNS)
    assert _rows(path, "Anmeldung")[1][:2] == ("Max", "Judo")
    assert _rows(path, "Abmeldung")[1][:2] == ("Eva", "Umzug")


def test_new_workbook_has_rule_columns_then_dates_then_mail_text(tmp_path: Path) -> None:
    path = tmp_path / "output.xlsx"

    _output(path, DATA_COLUMNS).save()

    workbook = load_workbook(path)
    assert workbook.sheetnames == ["daten", "fehler"]
    assert [cell.value for cell in workbook["daten"][1]] == [*DATA_COLUMNS, *FIXED_COLUMNS]
    assert [cell.value for cell in workbook["fehler"][1]] == list(ERROR_COLUMNS)
    assert workbook["daten"].freeze_panes == "A2"


def test_append_data_writes_values_dates_and_mail_text(tmp_path: Path) -> None:
    path = tmp_path / "output.xlsx"
    output = _output(path, DATA_COLUMNS)

    output.append_data(
        {"Name": "Jan", "Kurs": "Experimente", "Telefonnummer": "1234", "Zeit": "Mo", "Mail-Adresse": "max@mail.com"},
        "Mail text",
        received=RECEIVED_AT,
    )
    output.save()

    sheet = load_workbook(path)["daten"]
    values = [cell.value for cell in sheet[2]]
    assert values == ["max@mail.com", "Jan", "Experimente", "Mo", "1234", RECEIVED_AT, RUN_AT, "Mail text"]
    # Real Excel dates, so they sort and filter correctly.
    assert sheet["F2"].number_format == DATE_FORMAT
    assert sheet["G2"].number_format == DATE_FORMAT


def test_mail_without_usable_date_leaves_the_cell_empty(tmp_path: Path) -> None:
    path = tmp_path / "output.xlsx"
    output = _output(path, ["Name"])

    output.append_data({"Name": "Jan"}, "Text", received=None)
    output.save()

    assert _rows(path)[1] == ("Jan", None, RUN_AT, "Text")


def test_values_starting_with_equals_are_stored_as_text_not_formulas(tmp_path: Path) -> None:
    path = tmp_path / "output.xlsx"
    output = _output(path, ["Name"])
    payload = '=HYPERLINK("http://evil.example","klick")'

    output.append_data({"Name": payload}, payload)
    output.upsert_error(_error("k", reason=payload))
    output.save()

    workbook = load_workbook(path)
    assert workbook["daten"]["A2"].value == payload
    assert workbook["daten"]["A2"].data_type == "s"
    reason_cell = workbook["fehler"].cell(row=2, column=ERROR_COLUMNS.index("Grund") + 1)
    assert reason_cell.value == payload and reason_cell.data_type == "s"


def test_existing_workbook_is_appended_to(tmp_path: Path) -> None:
    path = tmp_path / "output.xlsx"
    first = _output(path, ["Name"])
    first.append_data({"Name": "A"}, "text A")
    first.save()

    second = _output(path, ["Name"])
    second.append_data({"Name": "B"}, "text B")
    second.save()

    assert [row[0] for row in _rows(path)] == ["Name", "A", "B"]


def test_reordered_rules_still_write_each_value_to_its_column(tmp_path: Path) -> None:
    path = tmp_path / "output.xlsx"
    _output(path, ["Name", "Kurs"]).save()

    output = _output(path, ["Kurs", "Name"])
    output.append_data({"Name": "Jan", "Kurs": "Experimente"}, "Text")
    output.save()

    rows = _rows(path)
    assert rows[0][:2] == ("Name", "Kurs")  # existing columns never move
    assert rows[1][:2] == ("Jan", "Experimente")


def test_new_rule_gets_a_new_column_at_the_end(tmp_path: Path) -> None:
    path = tmp_path / "output.xlsx"
    first = _output(path, ["Name"])
    first.append_data({"Name": "Alt"}, "alter Text")
    first.save()

    second = _output(path, ["Name", "Betreff"])
    second.append_data({"Name": "Neu", "Betreff": "Anmeldung"}, "neuer Text")
    second.save()

    rows = _rows(path)
    assert rows[0] == ("Name", *FIXED_COLUMNS, "Betreff")
    assert rows[1] == ("Alt", None, RUN_AT, "alter Text", None)
    assert rows[2] == ("Neu", None, RUN_AT, "neuer Text", "Anmeldung")


def test_own_columns_and_notes_in_excel_are_kept(tmp_path: Path) -> None:
    path = tmp_path / "output.xlsx"
    first = _output(path, ["Name"])
    first.append_data({"Name": "Anna"}, "Text A")
    first.save()
    workbook = load_workbook(path)
    sheet = workbook["daten"]
    sheet.cell(row=1, column=sheet.max_column + 1, value="Bestätigt")
    sheet.cell(row=2, column=sheet.max_column, value="ja")
    workbook.save(path)

    second = _output(path, ["Name"])
    second.append_data({"Name": "Ben"}, "Text B")
    second.save()

    rows = _rows(path)
    assert rows[0] == ("Name", *FIXED_COLUMNS, "Bestätigt")
    assert rows[1][-1] == "ja"
    assert rows[2] == ("Ben", None, RUN_AT, "Text B", None)


def test_removed_rule_keeps_its_column_and_new_rows_leave_it_empty(tmp_path: Path) -> None:
    path = tmp_path / "output.xlsx"
    first = _output(path, ["Name", "Kurs"])
    first.append_data({"Name": "Anna", "Kurs": "Chemie"}, "Text A")
    first.save()

    second = _output(path, ["Name"])
    second.append_data({"Name": "Ben"}, "Text B")
    second.save()

    rows = _rows(path)
    assert rows[0][:2] == ("Name", "Kurs")
    assert rows[1][:2] == ("Anna", "Chemie")
    assert rows[2][:2] == ("Ben", None)


def test_columns_added_next_to_unlabelled_notes_do_not_overwrite_them(tmp_path: Path) -> None:
    path = tmp_path / "output.xlsx"
    first = _output(path, ["Name"])
    first.append_data({"Name": "Anna"}, "Text")
    first.save()
    workbook = load_workbook(path)
    workbook["daten"].cell(row=2, column=6, value="Notiz ohne Überschrift")
    workbook.save(path)

    second = _output(path, ["Name", "Kurs"])
    second.save()

    rows = _rows(path)
    assert rows[1][5] == "Notiz ohne Überschrift"
    assert rows[0][6] == "Kurs"


def test_sheet_with_data_but_without_header_row_is_rejected(tmp_path: Path) -> None:
    path = tmp_path / "output.xlsx"
    workbook = Workbook()
    workbook.active.title = "daten"
    workbook.active.cell(row=2, column=1, value="Daten ohne Kopfzeile")
    workbook.save(path)

    with pytest.raises(SheetHeaderError, match="daten"):
        _output(path, ["Name"])


def test_workbook_from_before_the_date_and_text_columns_is_upgraded(tmp_path: Path) -> None:
    path = tmp_path / "output.xlsx"
    old = Workbook()
    old.active.title = "daten"
    old.active.append(["Name"])
    old.active.append(["Alt"])
    old.save(path)

    output = _output(path, ["Name"])
    output.append_data({"Name": "Neu"}, "Text")
    output.save()

    assert _rows(path) == [("Name", *FIXED_COLUMNS), ("Alt", None, None, None), ("Neu", None, RUN_AT, "Text")]


def test_mail_text_is_cleaned_and_truncated_to_excel_limit(tmp_path: Path) -> None:
    path = tmp_path / "output.xlsx"
    output = _output(path, ["Name"])

    output.append_data({"Name": "Jan\x01"}, "Zeile 1\x00\nZeile 2")
    output.append_data({"Name": "Lang"}, "x" * (EXCEL_CELL_LIMIT + 10))
    output.save()

    rows = _rows(path)[1:]
    assert (rows[0][0], rows[0][-1]) == ("Jan", "Zeile 1\nZeile 2")
    assert len(rows[1][-1]) == EXCEL_CELL_LIMIT
    assert rows[1][-1].endswith("gekürzt: Excel erlaubt höchstens 32.767 Zeichen pro Zelle]")


def test_check_writable_reports_locked_workbook(tmp_path: Path, monkeypatch) -> None:
    path = tmp_path / "output.xlsx"
    _output(path, ["Name"]).save()
    output = _output(path, ["Name"])

    def locked_open(self, *args, **kwargs):
        raise PermissionError("locked")

    monkeypatch.setattr(Path, "open", locked_open)

    with pytest.raises(OSError, match="open in Excel"):
        output.check_writable()


@pytest.mark.parametrize(
    "content",
    [
        b"",  # e.g. left behind by an interrupted copy
        bytes.fromhex("D0CF11E0A1B11AE1") + bytes(504),  # how Excel stores a workbook protected with a password
        b"Name;Kurs\n",  # a CSV file named .xlsx
    ],
)
def test_a_workbook_that_cannot_be_opened_is_reported_clearly(tmp_path: Path, content: bytes) -> None:
    path = tmp_path / "output.xlsx"
    path.write_bytes(content)

    with pytest.raises(WorkbookUnreadableError, match="output.xlsx"):
        _output(path, DATA_COLUMNS)
    assert path.read_bytes() == content  # left as it is


def test_sheet_names_differing_only_in_case_reuse_the_existing_sheet(tmp_path: Path) -> None:
    # Excel treats sheet names case-insensitively; e.g. a profile renamed from "anmeldung" to "Anmeldung",
    # or the data sheet setting changed from "daten" to "Daten", must keep the rows already there.
    path = tmp_path / "output.xlsx"
    first = ExcelOutput(path, {"anmeldung": ["Name"]}, "fehler", now=RUN_AT)
    first.append_data({"Name": "Anna"}, "Text")
    first.save()

    second = ExcelOutput(path, {"Anmeldung": ["Name"]}, "Fehler", now=RUN_AT)
    second.append_data({"Name": "Ben"}, "Text", sheet="Anmeldung")
    second.save()

    workbook = load_workbook(path)
    assert workbook.sheetnames == ["Anmeldung", "Fehler"]
    assert [row[0] for row in workbook["Anmeldung"].iter_rows(values_only=True)] == ["Name", "Anna", "Ben"]


def test_rows_and_columns_follow_the_data_not_formatted_empty_cells(tmp_path: Path) -> None:
    # Users format a range ahead in Excel (borders, number formats). Those cells are empty but count for
    # openpyxl's max_row/max_column; new rows must still follow the last row, new columns the last column.
    path = tmp_path / "output.xlsx"
    first = _output(path, ["Name"])
    first.append_data({"Name": "Anna"}, "Text")
    first.save()
    workbook = load_workbook(path)
    sheet = workbook["daten"]
    for row in sheet.iter_rows(min_row=3, max_row=200, max_col=30):
        for cell in row:
            cell.border = Border(bottom=Side(style="thin"))
    workbook.save(path)

    second = _output(path, ["Name", "Kurs"])
    second.append_data({"Name": "Ben", "Kurs": "Chemie"}, "Text")
    second.save()

    rows = _rows(path)
    assert [row[0] for row in rows[:3]] == ["Name", "Anna", "Ben"]
    header = [value for value in rows[0] if value is not None]
    assert header == ["Name", *FIXED_COLUMNS, "Kurs"]
    assert rows[0].index("Kurs") == len(header) - 1  # right after the used columns, not after column 30
