"""End-to-end proof that IMAP access is read-only and idempotent.

Runs the real `imaplib` against a scripted local IMAP server that logs every command and implements
the \\Seen side effect of non-PEEK fetches, then checks the mailbox is unchanged after two full runs.
"""

from __future__ import annotations

import base64
import copy
import imaplib
import re
import socketserver
import threading
from pathlib import Path

import pytest
from openpyxl import load_workbook

from mailprocessor.config import AppConfig, AppSection, FieldRule, ImapSourceConfig, ParsingRules, SourceConfig
from mailprocessor.processor import run_pipeline

ALLOWED_COMMANDS = {"CAPABILITY", "LOGIN", "AUTHENTICATE", "EXAMINE", "UID SEARCH", "UID FETCH", "LOGOUT"}


def _message(message_id: str | None, name: str) -> bytes:
    headers = [f"Message-ID: <{message_id}>"] if message_id else []
    headers += ["From: Form <form@example.com>", "Subject: Anmeldung", "Date: Mon, 05 Oct 2026 10:00:00 +0000"]
    return ("\r\n".join([*headers, "Content-Type: text/plain; charset=utf-8", "", f"Name: {name}", ""])).encode()


class MailboxState:
    def __init__(self) -> None:
        self.messages = {
            101: {"flags": set(), "body": _message("a@example.com", "Anna")},
            202: {"flags": {"\\Flagged"}, "body": _message(None, "Ben")},
        }
        self.commands: list[str] = []
        # (user, password) of each successful login, as the server decoded them.
        self.logins: list[tuple[str, str]] = []
        self.lock = threading.Lock()


class _Handler(socketserver.StreamRequestHandler):
    state: MailboxState

    def _send(self, line: str | bytes) -> None:
        self.wfile.write(line if isinstance(line, bytes) else line.encode())

    def handle(self) -> None:
        self._send("* OK [CAPABILITY IMAP4rev1 AUTH=PLAIN] scripted test server ready\r\n")
        while True:
            raw = self.rfile.readline()
            if not raw:
                return
            tag, _, rest = raw.decode().rstrip("\r\n").partition(" ")
            upper = rest.upper()
            verb = " ".join(upper.split()[:2]) if upper.startswith("UID ") else upper.split()[0]
            with self.state.lock:
                self.state.commands.append(rest)
            if verb == "CAPABILITY":
                self._send(f"* CAPABILITY IMAP4rev1\r\n{tag} OK done\r\n")
            elif verb == "LOGIN":
                _, user, password = rest.split(" ", 2)
                self.state.logins.append((user, password.strip('"')))
                self._send(f"{tag} OK logged in\r\n")
            elif verb == "AUTHENTICATE" and upper.split()[1:] == ["PLAIN"]:
                # RFC 4616: base64 of "authzid NUL user NUL password", UTF-8.
                self._send("+ \r\n")
                _authzid, user, password = base64.b64decode(self.rfile.readline().strip()).decode().split("\0")
                self.state.logins.append((user, password))
                self._send(f"{tag} OK authenticated\r\n")
            elif verb == "EXAMINE":
                self._send(f"* {len(self.state.messages)} EXISTS\r\n{tag} OK [READ-ONLY] EXAMINE completed\r\n")
            elif verb == "UID SEARCH":
                uids = " ".join(str(uid) for uid in sorted(self.state.messages))
                self._send(f"* SEARCH {uids}\r\n{tag} OK done\r\n")
            elif verb == "UID FETCH":
                uid = int(upper.split()[2])
                message = self.state.messages[uid]
                items = upper.split(" ", 3)[3]
                # RFC 3501: BODY[...] and RFC822 (but not their PEEK forms) implicitly set \Seen.
                if re.search(r"(?<!\.PEEK)BODY\[|RFC822(?![.]HEADER|[.]SIZE)", items):
                    message["flags"].add("\\Seen")
                body = message["body"]
                self._send(f"* 1 FETCH (UID {uid} BODY[] {{{len(body)}}}\r\n".encode() + body + b")\r\n")
                self._send(f"{tag} OK done\r\n")
            elif verb == "LOGOUT":
                self._send(f"* BYE\r\n{tag} OK bye\r\n")
                return
            else:
                self._send(f"{tag} BAD command not supported by test server\r\n")


@pytest.fixture
def imap_server(monkeypatch):
    state = MailboxState()
    handler = type("Handler", (_Handler,), {"state": state})
    server = socketserver.ThreadingTCPServer(("127.0.0.1", 0), handler)
    server.daemon_threads = True
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    # Plain imaplib against the local test server (TLS is covered by unit tests).
    monkeypatch.setattr(
        "mailprocessor.sources.imap_source.default_client_factory",
        lambda use_ssl: lambda host, port: imaplib.IMAP4(host, port),
    )
    try:
        yield state, server.server_address[1]
    finally:
        server.shutdown()
        server.server_close()


def _config(tmp_path: Path, port: int, mailbox: str = "INBOX", password: str = "pw") -> AppConfig:
    return AppConfig(
        app=AppSection(sqlite_path=str(tmp_path / "ledger.db"), output_xlsx=str(tmp_path / "out.xlsx")),
        source=SourceConfig(
            type="imap",
            imap=ImapSourceConfig(
                host="127.0.0.1", port=port, username="u", password=password, mailbox=mailbox, use_ssl=True
            ),
        ),
    )


def test_pipeline_runs_never_modify_the_mailbox_and_are_idempotent(imap_server, tmp_path: Path) -> None:
    state, port = imap_server
    original = copy.deepcopy(state.messages)
    rules = ParsingRules(fields=[FieldRule(column="Name", pattern=r"(?m)^Name:\s*(.+)$")])
    config = _config(tmp_path, port)

    first = run_pipeline(config, rules)
    second = run_pipeline(config, rules)

    assert (first.processed, first.skipped) == (2, 0)
    assert (second.processed, second.skipped) == (0, 2)
    rows = load_workbook(tmp_path / "out.xlsx")["daten"].iter_rows()
    assert [row[0].value for row in rows] == ["Name", "Anna", "Ben"]

    # Mailbox completely unchanged: same messages, same flags (no \Seen added), same content.
    assert state.messages == original

    verbs = {
        " ".join(cmd.upper().split()[:2]) if cmd.upper().startswith("UID ") else cmd.upper().split()[0]
        for cmd in state.commands
    }
    assert verbs <= ALLOWED_COMMANDS, verbs - ALLOWED_COMMANDS
    assert all("BODY.PEEK[]" in cmd for cmd in state.commands if cmd.upper().startswith("UID FETCH"))


def test_dry_run_reads_without_touching_mailbox_or_disk(imap_server, tmp_path: Path) -> None:
    state, port = imap_server
    original = copy.deepcopy(state.messages)
    config = _config(tmp_path, port)
    config.app.dry_run = True

    summary = run_pipeline(config, ParsingRules(fields=[FieldRule(column="Name", pattern=r"(?m)^Name:\s*(.+)$")]))

    assert summary.processed == 2
    assert state.messages == original
    assert list(tmp_path.iterdir()) == []


def test_mailbox_names_with_spaces_and_umlauts_reach_the_server_quoted(imap_server, tmp_path: Path) -> None:
    state, port = imap_server
    config = _config(tmp_path, port, mailbox="Anfragen Schüler")

    summary = run_pipeline(config, ParsingRules(fields=[FieldRule(column="Name", pattern=r"(?m)^Name:\s*(.+)$")]))

    assert summary.processed == 2
    assert 'EXAMINE "Anfragen Sch&APw-ler"' in state.commands


def test_passwords_with_non_ascii_characters_log_in(imap_server, tmp_path: Path) -> None:
    state, port = imap_server
    config = _config(tmp_path, port, password="Schlüssel§2026€")

    summary = run_pipeline(config, ParsingRules(fields=[FieldRule(column="Name", pattern=r"(?m)^Name:\s*(.+)$")]))

    assert summary.processed == 2
    assert state.logins == [("u", "Schlüssel§2026€")]


def test_ascii_passwords_still_use_login(imap_server, tmp_path: Path) -> None:
    state, port = imap_server

    rules = ParsingRules(fields=[FieldRule(column="Name", pattern=r"(?m)^Name:\s*(.+)$")])

    run_pipeline(_config(tmp_path, port), rules)

    assert state.logins == [("u", "pw")]
    assert not any(command.upper().startswith("AUTHENTICATE") for command in state.commands)
