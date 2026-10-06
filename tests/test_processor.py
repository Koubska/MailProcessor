"""run_pipeline: writing the workbook and the ledger, dry runs, and robustness against bad mails."""

import logging
import threading
from pathlib import Path

import pytest
from openpyxl import load_workbook
from pipeline_helpers import GOOD_BODY, _build_config, _build_rules, _data_rows, _snapshot, _write_eml

import mailprocessor.processor as processor
import mailprocessor.sources.eml_folder_source as eml_source
from mailprocessor.errors import RunInProgressError
from mailprocessor.excel_writer import ERROR_COLUMNS
from mailprocessor.processor import RunSummary, run_pipeline


def test_run_pipeline_writes_data_and_error_sheets(tmp_path: Path) -> None:
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    _write_eml(
        inbox / "good.eml",
        [
            "Von: Max Mustermann <max.mustermann@mail.com>",
            "hiermit möchte ich mein Kind Jan Must+ für folgendes Angebot anmelden:",
            "Tag: Mo, 23.11.2026 (14:15-16:00 Uhr)",
            "Angebot: Experimente",
            "Telefonnummer: 1234 567890",
        ],
        "good@example.com",
    )
    _write_eml(
        inbox / "bad.eml",
        [
            "Von: Max Mustermann <max.mustermann@mail.com>",
            "hiermit möchte ich mein Kind Jan Must+ für folgendes Angebot anmelden:",
            "Tag: Mo, 23.11.2026 (14:15-16:00 Uhr)",
            "Angebot: Experimente",
        ],
        "bad@example.com",
    )

    summary = run_pipeline(_build_config(tmp_path), _build_rules())

    assert summary.seen == 2
    assert summary.processed == 1
    assert summary.failed == 1
    assert summary.skipped == 0

    workbook = load_workbook(tmp_path / "out" / "mail_export.xlsx")
    assert workbook["daten"].max_row == 2
    assert workbook["fehler"].max_row == 2


def test_run_pipeline_is_idempotent_for_successful_messages(tmp_path: Path) -> None:
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    _write_eml(
        inbox / "good.eml",
        [
            "Von: Max Mustermann <max.mustermann@mail.com>",
            "hiermit möchte ich mein Kind Jan Must+ für folgendes Angebot anmelden:",
            "Tag: Mo, 23.11.2026 (14:15-16:00 Uhr)",
            "Angebot: Experimente",
            "Telefonnummer: 1234 567890",
        ],
        "good@example.com",
    )
    config = _build_config(tmp_path)
    rules = _build_rules()

    first = run_pipeline(config, rules)
    second = run_pipeline(config, rules)

    assert first.processed == 1
    assert second.skipped == 1
    assert second.processed == 0

    workbook = load_workbook(tmp_path / "out" / "mail_export.xlsx")
    assert workbook["daten"].max_row == 2


def test_dry_run_writes_neither_workbook_nor_ledger(tmp_path: Path) -> None:
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    _write_eml(inbox / "good.eml", GOOD_BODY, "good@example.com")
    config = _build_config(tmp_path)
    config.app.dry_run = True

    summary = run_pipeline(config, _build_rules())

    assert summary.processed == 1
    assert not (tmp_path / "out").exists()
    assert not (tmp_path / "data").exists()


def test_unreadable_message_does_not_abort_batch(tmp_path: Path) -> None:
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    _write_eml(inbox / "a-good.eml", GOOD_BODY, "good@example.com")
    (inbox / "b-bogus-charset.eml").write_bytes(
        b"Message-ID: <bogus@example.com>\nContent-Type: text/plain; charset=x-unknown-8bit\n\n"
        + "\n".join(GOOD_BODY).encode("utf-8")
    )
    _write_eml(inbox / "c-good.eml", GOOD_BODY, "good2@example.com")

    summary = run_pipeline(_build_config(tmp_path), _build_rules())

    assert summary.processed == 3
    assert summary.failed == 0


def test_eml_read_errors_are_reported_and_batch_continues(tmp_path: Path, monkeypatch) -> None:
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    _write_eml(inbox / "a-broken.eml", GOOD_BODY, "broken@example.com")
    _write_eml(inbox / "b-good.eml", GOOD_BODY, "good@example.com")

    original = eml_source.parse_eml_file

    def flaky_parse(path: Path, source_location: str):
        if path.name == "a-broken.eml":
            raise OSError("disk error")
        return original(path, source_location)

    monkeypatch.setattr(eml_source, "parse_eml_file", flaky_parse)

    summary = run_pipeline(_build_config(tmp_path), _build_rules())

    assert (summary.processed, summary.failed) == (1, 1)
    errors = load_workbook(tmp_path / "out" / "mail_export.xlsx")["fehler"]
    row = dict(zip(ERROR_COLUMNS, next(errors.iter_rows(min_row=2, values_only=True)), strict=True))
    assert row["E-Mail"] == "a-broken.eml"
    assert row["Grund"].startswith("E-Mail konnte nicht gelesen werden")
    assert "disk error" in row["Grund"]


def test_max_messages_counts_only_new_messages(tmp_path: Path) -> None:
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    for index in range(5):
        _write_eml(inbox / f"mail-{index}.eml", GOOD_BODY, f"m{index}@example.com")
    config = _build_config(tmp_path, max_messages=2)
    rules = _build_rules()

    results = [run_pipeline(config, rules) for _ in range(3)]

    assert [r.processed for r in results] == [2, 2, 1]
    assert load_workbook(tmp_path / "out" / "mail_export.xlsx")["daten"].max_row == 6


def test_locked_workbook_leaves_ledger_untouched(tmp_path: Path, monkeypatch) -> None:
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    _write_eml(inbox / "mail.eml", GOOD_BODY, "m@example.com")
    config = _build_config(tmp_path)

    def failing_save(self) -> None:
        raise OSError("locked")

    monkeypatch.setattr("mailprocessor.excel_writer.ExcelOutput.save", failing_save)
    with pytest.raises(OSError, match="locked"):
        run_pipeline(config, _build_rules())
    monkeypatch.undo()

    summary = run_pipeline(config, _build_rules())

    assert summary.processed == 1
    assert load_workbook(tmp_path / "out" / "mail_export.xlsx")["daten"].max_row == 2


def test_missing_eml_folder_is_an_error(tmp_path: Path) -> None:
    with pytest.raises(FileNotFoundError, match="EML folder does not exist"):
        run_pipeline(_build_config(tmp_path), _build_rules())


def test_eml_files_remain_completely_unchanged(tmp_path: Path) -> None:
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    _write_eml(inbox / "good.eml", GOOD_BODY, "good@example.com")
    _write_eml(inbox / "bad.eml", GOOD_BODY[:2], "bad@example.com")
    before = _snapshot(inbox)

    run_pipeline(_build_config(tmp_path), _build_rules())
    run_pipeline(_build_config(tmp_path), _build_rules())

    assert _snapshot(inbox) == before


def test_moved_mail_folder_is_not_exported_again(tmp_path: Path) -> None:
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    _write_eml(inbox / "good.eml", GOOD_BODY, "good@example.com")
    (inbox / "no-id.eml").write_text("From: a@example.com\n\n" + "\n".join(GOOD_BODY), encoding="utf-8")
    config = _build_config(tmp_path)
    assert run_pipeline(config, _build_rules()).processed == 2

    inbox.rename(tmp_path / "moved")
    assert config.source.eml is not None
    config.source.eml.folder = str(tmp_path / "moved")
    second = run_pipeline(config, _build_rules())

    assert (second.processed, second.skipped) == (0, 2)
    assert load_workbook(tmp_path / "out" / "mail_export.xlsx")["daten"].max_row == 3


def test_outputs_inside_mail_folder_are_rejected(tmp_path: Path) -> None:
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    config = _build_config(tmp_path)
    config.app.output_xlsx = str(inbox / "export.xlsx")

    with pytest.raises(ValueError, match="must not be inside the mail folder"):
        run_pipeline(config, _build_rules())
    assert list(inbox.iterdir()) == []


def test_run_logs_progress_and_failure_reasons_without_mail_content(tmp_path: Path, caplog) -> None:
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    _write_eml(inbox / "good.eml", GOOD_BODY, "good@example.com")
    _write_eml(inbox / "missing-phone.eml", GOOD_BODY[:4], "bad@example.com")
    config = _build_config(tmp_path)

    with caplog.at_level(logging.DEBUG, logger="mailprocessor"):
        run_pipeline(config, _build_rules())
        run_pipeline(config, _build_rules())

    messages = [record.getMessage() for record in caplog.records]
    assert any(m.startswith("Found 2 file(s) matching *.eml") for m in messages)
    assert "Processed good.eml (<good@example.com>)" in messages
    assert any(
        r.levelno == logging.WARNING
        and r.getMessage() == "Failed missing-phone.eml (<bad@example.com>): Required fields missing: Telefonnummer"
        for r in caplog.records
    )
    assert "Skipped good.eml (<good@example.com>): already processed in an earlier run" in messages
    assert any("found Mail-Adresse, Name, Kurs, Zeit; missing Telefonnummer" in m for m in messages)
    assert any(m.startswith("Run finished: seen=2 processed=0 skipped=1 failed=1") for m in messages)
    # Confidential content (extracted values, body text) never appears in logs.
    log_text = "\n".join(messages)
    assert "Experimente" not in log_text
    assert "1234 567890" not in log_text


def test_info_level_hides_per_field_details(tmp_path: Path, caplog) -> None:
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    _write_eml(inbox / "good.eml", GOOD_BODY, "good@example.com")

    with caplog.at_level(logging.INFO, logger="mailprocessor"):
        run_pipeline(_build_config(tmp_path), _build_rules())

    messages = [record.getMessage() for record in caplog.records]
    assert "Processed good.eml (<good@example.com>)" in messages
    assert not any("found " in m for m in messages)


def test_data_rows_include_the_full_mail_text(tmp_path: Path) -> None:
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    _write_eml(inbox / "good.eml", GOOD_BODY, "good@example.com")

    run_pipeline(_build_config(tmp_path), _build_rules())

    sheet = load_workbook(tmp_path / "out" / "mail_export.xlsx")["daten"]
    assert sheet.cell(row=1, column=8).value == "E-Mail-Inhalt"
    assert sheet.cell(row=2, column=8).value == "\n".join(GOOD_BODY)


def test_control_characters_in_a_mail_do_not_abort_the_run(tmp_path: Path) -> None:
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    body = [line.replace("Jan Must+", "Jan\x01 Must+") for line in GOOD_BODY]
    _write_eml(inbox / "control-char.eml", body, "ctrl@example.com")

    summary = run_pipeline(_build_config(tmp_path), _build_rules())

    assert summary.processed == 1
    assert load_workbook(tmp_path / "out" / "mail_export.xlsx")["daten"]["B2"].value == "Jan Must+"


def test_overlapping_runs_never_lose_rows(tmp_path: Path, monkeypatch) -> None:
    # Run B has loaded the workbook and waits for its mail source (e.g. a slow IMAP login) while run A
    # (a second window, or a scheduled CLI run) exports the mail. B must not save its older copy over A's rows.
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    _write_eml(inbox / "x.eml", GOOD_BODY, "x@example.com")
    config = _build_config(tmp_path)
    b_waits, release_b = threading.Event(), threading.Event()
    read_mails = processor.iter_eml_messages

    def slow_for_b(*args, **kwargs):
        if threading.current_thread().name == "B":
            b_waits.set()
            release_b.wait(10)
        yield from read_mails(*args, **kwargs)

    monkeypatch.setattr(processor, "iter_eml_messages", slow_for_b)
    monkeypatch.setattr("mailprocessor.ledger.LOCK_TIMEOUT_SECONDS", 0.1)
    outcomes: dict[str, object] = {}

    def run(name: str) -> None:
        try:
            outcomes[name] = run_pipeline(config, _build_rules())
        except Exception as exc:
            outcomes[name] = exc

    run_b = threading.Thread(target=run, args=("B",), name="B")
    run_b.start()
    assert b_waits.wait(10)
    run("A")
    release_b.set()
    run_b.join(10)

    assert [row["Name"] for row in _data_rows(tmp_path)] == ["Jan Must+"]
    # The run that started second is refused with a clear message instead of racing the first.
    assert isinstance(outcomes["A"], RunInProgressError)
    assert isinstance(outcomes["B"], RunSummary) and outcomes["B"].processed == 1
