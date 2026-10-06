"""The subject/sender filter: mails it leaves out are skipped without an error."""

from pathlib import Path

from openpyxl import load_workbook
from pipeline_helpers import GOOD_BODY, _build_config, _build_rules, _write_eml

from mailprocessor.config import (
    MailFilter,
)
from mailprocessor.models import MailReadError
from mailprocessor.processor import run_pipeline


def test_mails_not_matching_the_filter_are_left_out_without_an_error(tmp_path: Path) -> None:
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    _write_eml(inbox / "form.eml", GOOD_BODY, "form@example.com", subject="Kontaktformular: Anmeldung")
    _write_eml(inbox / "news.eml", ["Unser Newsletter im Oktober"], "news@example.com", subject="Newsletter")
    config = _build_config(tmp_path)
    config.filter = MailFilter(subject=["kontaktformular"])
    calls: list[tuple[int, int]] = []

    summary = run_pipeline(config, _build_rules(), progress=lambda current, total: calls.append((current, total)))

    assert (summary.seen, summary.processed, summary.failed, summary.filtered) == (1, 1, 0, 1)
    assert summary.problems == ()
    assert calls[-1] == (2, 2)
    workbook = load_workbook(tmp_path / "out" / "mail_export.xlsx")
    assert workbook["daten"].max_row == 2
    assert workbook["fehler"].max_row == 1


def test_mails_left_out_by_the_filter_are_read_once_the_filter_changes(tmp_path: Path) -> None:
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    _write_eml(inbox / "other.eml", GOOD_BODY, "other@example.com", subject="Rückfrage")
    config = _build_config(tmp_path)
    config.filter = MailFilter(subject=["Kontaktformular"])
    assert run_pipeline(config, _build_rules()).filtered == 1

    config.filter = MailFilter()
    summary = run_pipeline(config, _build_rules())

    assert (summary.processed, summary.filtered) == (1, 0)


def test_unreadable_mails_are_reported_even_with_a_filter(tmp_path: Path, monkeypatch) -> None:
    (tmp_path / "inbox").mkdir()
    config = _build_config(tmp_path)
    config.filter = MailFilter(subject=["Kontaktformular"])
    unreadable = MailReadError(source_type="eml", source_location="x", message_identity="file:x", reason="kaputt")

    def source(*_args, **_kwargs):
        yield unreadable

    monkeypatch.setattr("mailprocessor.processor.iter_source_messages", source)

    summary = run_pipeline(config, _build_rules())

    # Subject and sender of an unreadable mail are unknown, so it may well be one of the wanted ones.
    assert (summary.failed, summary.filtered) == (1, 0)


def test_setting_a_filter_clears_old_error_rows_of_mails_it_leaves_out(tmp_path: Path) -> None:
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    _write_eml(inbox / "news.eml", ["Unser Newsletter im Oktober"], "news@example.com", subject="Newsletter")
    config = _build_config(tmp_path)
    assert run_pipeline(config, _build_rules()).failed == 1

    config.filter = MailFilter(subject=["Kontaktformular"])
    run_pipeline(config, _build_rules())

    assert load_workbook(tmp_path / "out" / "mail_export.xlsx")["fehler"].max_row == 1
