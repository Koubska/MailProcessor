"""Several profiles: one shared sheet with a profile column, or one sheet per profile."""

from pathlib import Path

import pytest
from openpyxl import load_workbook
from pipeline_helpers import GOOD_BODY, _build_config, _build_rules, _write_eml

from mailprocessor.config import (
    FieldRule,
    ParsingRules,
    Profile,
)
from mailprocessor.processor import run_pipeline

ABMELDUNG_BODY = [
    "Von: Eva Muster <eva@example.com>",
    "hiermit melde ich mein Kind ab.",
    "Angebot: Experimente",
    "Grund: Umzug",
]


def _two_profiles() -> ParsingRules:
    return ParsingRules(
        profiles=[
            Profile(name="Anmeldung", fields=_build_rules().profiles[0].fields),
            Profile(
                name="Abmeldung",
                fields=[
                    FieldRule(column="Mail-Adresse", type="email", label=["Von", "From"]),
                    FieldRule(column="Kurs", type="label", label="Angebot:"),
                    FieldRule(column="Grund", type="label", label="Grund:"),
                ],
            ),
        ]
    )


def _write_profile_mails(inbox: Path) -> None:
    inbox.mkdir()
    _write_eml(inbox / "1_anmeldung.eml", GOOD_BODY, "an@example.com")
    _write_eml(inbox / "2_abmeldung.eml", ABMELDUNG_BODY, "ab@example.com")
    _write_eml(inbox / "3_unklar.eml", ["Grund: keiner"], "unklar@example.com")


def test_profiles_share_one_sheet_with_a_profile_column(tmp_path: Path) -> None:
    _write_profile_mails(tmp_path / "inbox")

    summary = run_pipeline(_build_config(tmp_path), _two_profiles())

    assert (summary.processed, summary.failed) == (2, 1)
    assert summary.per_profile == (("Anmeldung", 1), ("Abmeldung", 1))
    workbook = load_workbook(tmp_path / "out" / "mail_export.xlsx")
    assert workbook.sheetnames == ["daten", "fehler"]
    rows = list(workbook["daten"].iter_rows(values_only=True))
    assert rows[0][:7] == ("Profil", "Mail-Adresse", "Name", "Kurs", "Zeit", "Telefonnummer", "Grund")
    assert rows[1][:4] == ("Anmeldung", "max.mustermann@mail.com", "Jan Must+", "Experimente")
    assert rows[2][:7] == ("Abmeldung", "eva@example.com", None, "Experimente", None, None, "Umzug")


def test_profiles_can_write_one_sheet_each(tmp_path: Path) -> None:
    _write_profile_mails(tmp_path / "inbox")
    config = _build_config(tmp_path)
    config.app.profile_sheets = "per_profile"

    run_pipeline(config, _two_profiles())

    workbook = load_workbook(tmp_path / "out" / "mail_export.xlsx")
    assert workbook.sheetnames == ["Anmeldung", "Abmeldung", "fehler"]
    anmeldung = list(workbook["Anmeldung"].iter_rows(values_only=True))
    abmeldung = list(workbook["Abmeldung"].iter_rows(values_only=True))
    assert len(anmeldung) == len(abmeldung) == 2
    assert abmeldung[0][:3] == ("Mail-Adresse", "Kurs", "Grund")
    assert abmeldung[1][:3] == ("eva@example.com", "Experimente", "Umzug")


def test_failed_mail_names_the_closest_profile(tmp_path: Path) -> None:
    _write_profile_mails(tmp_path / "inbox")

    summary = run_pipeline(_build_config(tmp_path), _two_profiles())

    # "Grund: keiner" + the From header fit Abmeldung best; only its "Kurs" is missing.
    (problem,) = summary.problems
    assert (problem.profile, problem.missing) == ("Abmeldung", ("Kurs",))
    errors = list(load_workbook(tmp_path / "out" / "mail_export.xlsx")["fehler"].iter_rows(values_only=True))
    row = dict(zip(errors[0], errors[1], strict=False))
    assert row["Fehlende Felder"] == "Kurs"
    assert row["Grund"] == "Pflichtfelder nicht gefunden (am ähnlichsten: Profil „Abmeldung“)"


def test_single_profile_has_no_profile_column(tmp_path: Path) -> None:
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    _write_eml(inbox / "good.eml", GOOD_BODY, "good@example.com")

    summary = run_pipeline(_build_config(tmp_path), _build_rules())

    header = next(load_workbook(tmp_path / "out" / "mail_export.xlsx")["daten"].iter_rows(values_only=True))
    assert "Profil" not in header
    assert summary.per_profile == ()


def test_profile_sheet_must_not_be_the_error_sheet(tmp_path: Path) -> None:
    (tmp_path / "inbox").mkdir()
    config = _build_config(tmp_path)
    config.app.profile_sheets = "per_profile"
    rules = ParsingRules(profiles=[Profile(name="Fehler", fields=_build_rules().profiles[0].fields)])

    with pytest.raises(ValueError, match="same name as the sheet for problems"):
        run_pipeline(config, rules)
