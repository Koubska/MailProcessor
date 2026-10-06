"""Progress, stopping a run, and starting over ("Alles neu exportieren")."""

import threading
from datetime import UTC, datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook
from pipeline_helpers import (
    GOOD_BODY,
    _build_config,
    _build_rules,
    _inbox_with_good_mails,
    _snapshot,
    _write_eml_with_date,
)

from mailprocessor.errors import LedgerUnreadableError
from mailprocessor.ledger import Ledger
from mailprocessor.processor import run_pipeline, start_over


def test_progress_reports_each_message_with_total(tmp_path: Path) -> None:
    inbox = _inbox_with_good_mails(tmp_path, 3)
    _write_eml_with_date(inbox / "old.eml", GOOD_BODY, "old@example.com", "Mon, 01 Jan 2001 10:00:00 +0000")
    config = _build_config(tmp_path)
    config.app.max_age_days = 30
    calls: list[tuple[int, int]] = []

    summary = run_pipeline(config, _build_rules(), progress=lambda current, total: calls.append((current, total)))

    # The old mail is filtered out but still counts, so the progress reaches the total.
    assert calls == [(1, 4), (2, 4), (3, 4), (4, 4)]
    assert summary.processed == 3


def test_cancel_saves_what_was_processed_and_next_run_continues(tmp_path: Path) -> None:
    _inbox_with_good_mails(tmp_path, 5)
    config = _build_config(tmp_path)
    cancel = threading.Event()

    def stop_after_two(current: int, _total: int) -> None:
        if current == 3:  # mails 1 and 2 are done when the third one is read
            cancel.set()

    first = run_pipeline(config, _build_rules(), progress=stop_after_two, cancel=cancel)
    assert (first.processed, first.cancelled) == (2, True)
    assert load_workbook(tmp_path / "out" / "mail_export.xlsx")["daten"].max_row == 3

    second = run_pipeline(config, _build_rules())
    assert (second.processed, second.skipped, second.cancelled) == (3, 2, False)
    assert load_workbook(tmp_path / "out" / "mail_export.xlsx")["daten"].max_row == 6


def test_start_over_keeps_backup_and_exports_everything_again(tmp_path: Path) -> None:
    inbox = _inbox_with_good_mails(tmp_path, 2)
    config = _build_config(tmp_path)
    run_pipeline(config, _build_rules())
    before = _snapshot(inbox)

    backup = start_over(config, now=datetime(2026, 10, 5, 21, 30, 0, tzinfo=UTC))

    assert backup == tmp_path / "out" / "mail_export_backup_2026-10-05_21-30-00.xlsx"
    assert load_workbook(backup)["daten"].max_row == 3
    assert not (tmp_path / "out" / "mail_export.xlsx").exists()
    with Ledger(tmp_path / "data" / "ledger.db", read_only=True) as ledger:
        assert ledger.count_processed() == 0
    assert _snapshot(inbox) == before

    again = run_pipeline(config, _build_rules())
    assert (again.processed, again.skipped) == (2, 0)
    assert load_workbook(tmp_path / "out" / "mail_export.xlsx")["daten"].max_row == 3


def test_start_over_replaces_a_damaged_ledger(tmp_path: Path) -> None:
    # A damaged ledger stops every run; "Alles neu exportieren" must still work as the way out.
    _inbox_with_good_mails(tmp_path, 2)
    config = _build_config(tmp_path)
    ledger_path = tmp_path / "data" / "ledger.db"
    ledger_path.parent.mkdir()
    ledger_path.write_bytes(b"not a database" * 100)
    with pytest.raises(LedgerUnreadableError):
        run_pipeline(config, _build_rules())

    start_over(config, now=datetime(2026, 10, 5, 21, 30, 0, tzinfo=UTC))

    assert (tmp_path / "data" / "ledger_damaged_2026-10-05_21-30-00.db").read_bytes() == b"not a database" * 100
    again = run_pipeline(config, _build_rules())
    assert (again.processed, again.skipped) == (2, 0)


def test_start_over_without_previous_run_does_nothing(tmp_path: Path) -> None:
    assert start_over(_build_config(tmp_path)) is None
    assert not (tmp_path / "data").exists()
